from __future__ import annotations

import copy
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
ILP = ROOT / "benchmarks/eplb_rebalance/ILP"
E2E = ROOT / ("benchmarks/eplb_rebalance/results/serving_nixl_20260911_async_random500")
PROFILE = ROOT / (
    "benchmarks/eplb_rebalance/results/serving_nixl_20260911_async_random500_profile"
)
THESIS_IMAGES = ROOT / "benchmarks/eplb_rebalance/thesis/Images"
OUT = HERE / "eplb_migration_batching.pptx"
LOGO = HERE / "polimi_logo.png"
PR_URL = "https://github.com/vllm-project/vllm/pull/52641"
OMNI_URL = "https://github.com/omni-ai-npu/omni-infer"


# Restrained academic palette, adapted from the reference presentation.
NAVY = RGBColor(19, 48, 82)
BLUE = RGBColor(45, 103, 164)
LIGHT_BLUE = RGBColor(225, 237, 248)
CYAN = RGBColor(21, 142, 166)
GREEN = RGBColor(27, 141, 96)
LIGHT_GREEN = RGBColor(229, 245, 238)
ORANGE = RGBColor(221, 133, 42)
LIGHT_ORANGE = RGBColor(251, 239, 222)
RED = RGBColor(191, 65, 73)
LIGHT_RED = RGBColor(250, 230, 232)
PURPLE = RGBColor(129, 83, 160)
INK = RGBColor(28, 39, 53)
MUTED = RGBColor(93, 108, 126)
GRID = RGBColor(210, 220, 231)
BG = RGBColor(249, 251, 253)
WHITE = RGBColor(255, 255, 255)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def ensure_logo() -> None:
    if LOGO.exists():
        return
    source = THESIS_IMAGES / (
        "02_ING_IND_INFORMAZIONE_COL_POSITIVO-eps-converted-to.pdf"
    )
    subprocess.run(
        [
            "pdftoppm",
            "-png",
            "-singlefile",
            "-r",
            "220",
            str(source),
            str(LOGO.with_suffix("")),
        ],
        check=True,
    )


def load_equations() -> dict[str, etree._Element]:
    formulas = {
        "assign": r"\sum_{b\in B}x_{p,b}=1,\qquad \forall p\in P",
        "endpoint": (
            r"\sum_{p:s_p=r\,\lor\,d_p=r}x_{p,b}\le 1,"
            r"\qquad \forall r\in R,\;b\in B"
        ),
        "objective": (
            r"\min T_{\mathrm{infer}},\qquad T_{\mathrm{infer}}\ge "
            r"T_0\left(1+\frac{Q_{l,b}}{C_l(1-\rho)}\right),"
            r"\quad \forall l,b"
        ),
        "mask": (
            r"m_{s,d}=(1\ll s)\lor(1\ll d),\qquad "
            r"m_{s,d}\land u_b=0"
        ),
        "complexity": (r"O(S+M+FK)\ \mathrm{time},\qquad O(F)\ \mathrm{space}"),
        "ratio": (
            r"R=\frac{C_{\max}/1000}{T_{\mathrm{off}}-T_{\mathrm{on}}}"
            r"\times100\%=0.285\%"
        ),
    }
    markdown = []
    for key, formula in formulas.items():
        markdown.extend([f"# {key}", "", f"$${formula}$$", "", "---", ""])

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        source = tmp_path / "equations.md"
        deck = tmp_path / "equations.pptx"
        source.write_text("\n".join(markdown), encoding="utf-8")
        subprocess.run(["pandoc", str(source), "-o", str(deck)], check=True)
        result: dict[str, etree._Element] = {}
        ns = {
            "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
        }
        with zipfile.ZipFile(deck) as archive:
            for index, key in enumerate(formulas, start=1):
                xml = archive.read(f"ppt/slides/slide{index}.xml")
                root = etree.fromstring(xml)
                math = root.find(".//a14:m", namespaces=ns)
                if math is None:
                    raise RuntimeError(f"Pandoc did not generate Office Math: {key}")
                result[key] = copy.deepcopy(math)
    return result


ensure_logo()
EQUATIONS = load_equations()


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_box(
    slide,
    x,
    y,
    w,
    h,
    *,
    fill=WHITE,
    line=GRID,
    radius=True,
    width=1,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(width)
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=20,
    color=INK,
    bold=False,
    font="Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin=0.04,
    url=None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if url:
        run.hyperlink.address = url
    return box


def add_bullets(slide, items, x, y, w, h, *, size=17, color=INK, spacing=9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.word_wrap = True
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.text_frame = frame if hasattr(p, "text_frame") else None
        p._p.get_or_add_pPr().insert(
            0,
            etree.Element(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar",
                char="•",
            ),
        )
    return box


def add_rich(slide, runs, x, y, w, h, *, size=18, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    for text, color, bold in runs:
        run = p.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_line(slide, x1, y1, x2, y2, *, color=GRID, width=2, arrow=False):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if arrow:
        line.line.end_arrowhead = True
    return line


def add_image_fit(slide, path, x, y, w, h, *, crop=False):
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = w / h
    if crop:
        picture = slide.shapes.add_picture(
            str(path), Inches(x), Inches(y), Inches(w), Inches(h)
        )
        if image_ratio > box_ratio:
            shown = box_ratio / image_ratio
            picture.crop_left = picture.crop_right = (1 - shown) / 2
        else:
            shown = image_ratio / box_ratio
            picture.crop_top = picture.crop_bottom = (1 - shown) / 2
        return picture
    if image_ratio > box_ratio:
        fitted_w = w
        fitted_h = w / image_ratio
        x0, y0 = x, y + (h - fitted_h) / 2
    else:
        fitted_h = h
        fitted_w = h * image_ratio
        x0, y0 = x + (w - fitted_w) / 2, y
    return slide.shapes.add_picture(
        str(path), Inches(x0), Inches(y0), Inches(fitted_w), Inches(fitted_h)
    )


def add_equation(slide, key, x, y, w, h, *, size=18, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tx_body = box._element.txBody
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    m_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    for paragraph in list(tx_body.findall(f"{{{a_ns}}}p")):
        tx_body.remove(paragraph)
    paragraph = etree.Element(f"{{{a_ns}}}p")
    p_pr = etree.SubElement(paragraph, f"{{{a_ns}}}pPr")
    p_pr.set("algn", "ctr")
    math = copy.deepcopy(EQUATIONS[key])
    for run in math.findall(f".//{{{m_ns}}}r"):
        props = etree.Element(f"{{{a_ns}}}rPr")
        props.set("lang", "en-US")
        props.set("sz", str(size * 100))
        props.set("dirty", "0")
        fill = etree.SubElement(props, f"{{{a_ns}}}solidFill")
        rgb = etree.SubElement(fill, f"{{{a_ns}}}srgbClr")
        rgb.set("val", "".join(f"{channel:02X}" for channel in color))
        latin = etree.SubElement(props, f"{{{a_ns}}}latin")
        latin.set("typeface", "Cambria Math")
        text = run.find(f"{{{m_ns}}}t")
        insertion = run.index(text) if text is not None else len(run)
        run.insert(insertion, props)
    paragraph.append(math)
    tx_body.append(paragraph)
    body_pr = tx_body.find(f"{{{a_ns}}}bodyPr")
    body_pr.set("anchor", "ctr")
    body_pr.set("wrap", "none")
    return box


def add_header(slide, title, number, section, subtitle=None):
    add_text(
        slide, section.upper(), 0.55, 0.18, 4.6, 0.25, size=9, color=BLUE, bold=True
    )
    add_text(slide, title, 0.55, 0.48, 12.1, 0.55, size=26, color=NAVY, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.57, 1.02, 12.0, 0.32, size=11, color=MUTED)
    add_line(slide, 0.58, 1.4, 12.76, 1.4, color=GRID, width=1)
    add_text(
        slide,
        f"{number:02d}",
        12.3,
        7.12,
        0.48,
        0.2,
        size=8,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )
    add_text(
        slide, "WU CAO · MSc Thesis Defense", 0.56, 7.12, 3.2, 0.2, size=8, color=MUTED
    )


def add_source(slide, text, url=None):
    add_text(
        slide,
        text,
        7.25,
        7.1,
        4.7,
        0.22,
        size=7,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        url=url,
    )


def add_metric(slide, value, label, x, y, w, *, color=BLUE, h=1.0, value_size=22):
    add_box(slide, x, y, w, h, fill=WHITE, line=color)
    add_text(
        slide,
        value,
        x + 0.1,
        y + 0.1,
        w - 0.2,
        0.38,
        size=value_size,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        label,
        x + 0.1,
        y + 0.55,
        w - 0.2,
        0.25,
        size=9,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def node(slide, label, x, y, *, color=BLUE, size=0.65):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = WHITE
    shape.line.width = Pt(1.5)
    add_text(
        slide,
        label,
        x,
        y + 0.01,
        size,
        size - 0.02,
        size=11,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    return shape


def section_marker(slide, labels, active):
    total_w = 11.8
    x0 = 0.76
    gap = 0.08
    width = (total_w - gap * (len(labels) - 1)) / len(labels)
    for index, label in enumerate(labels):
        color = BLUE if index == active else GRID
        text_color = WHITE if index == active else MUTED
        x = x0 + index * (width + gap)
        add_box(slide, x, 6.4, width, 0.42, fill=color, line=color, radius=False)
        add_text(
            slide,
            label,
            x + 0.03,
            6.47,
            width - 0.06,
            0.22,
            size=8,
            color=text_color,
            bold=index == active,
            align=PP_ALIGN.CENTER,
        )


def add_table(slide, rows, col_widths, x, y, row_h, *, font_size=11):
    colors = [NAVY] + [WHITE] * (len(rows) - 1)
    y0 = y
    for row_index, row in enumerate(rows):
        x0 = x
        for col_index, (value, width) in enumerate(zip(row, col_widths)):
            fill = colors[row_index]
            if row_index > 0 and row_index % 2 == 0:
                fill = RGBColor(243, 247, 251)
            add_box(slide, x0, y0, width, row_h, fill=fill, line=GRID, radius=False)
            add_text(
                slide,
                str(value),
                x0 + 0.04,
                y0 + 0.04,
                width - 0.08,
                row_h - 0.08,
                size=font_size if row_index else font_size - 1,
                color=WHITE if row_index == 0 else INK,
                bold=row_index == 0 or col_index == 0,
                align=PP_ALIGN.CENTER,
            )
            x0 += width
        y0 += row_h


# 01 — title
slide = prs.slides.add_slide(BLANK)
add_bg(slide, WHITE)
add_box(slide, 0, 0, 4.35, 7.5, fill=NAVY, line=NAVY, radius=False)
add_text(
    slide,
    "MASTER'S THESIS DEFENSE",
    0.63,
    0.57,
    3.1,
    0.32,
    size=11,
    color=RGBColor(163, 203, 232),
    bold=True,
)
add_text(
    slide,
    "Contention-Aware\nExpert Migration\nfor MoE Serving",
    0.63,
    1.2,
    3.58,
    2.4,
    size=27,
    color=WHITE,
    bold=True,
    valign=MSO_ANCHOR.TOP,
)
add_text(
    slide,
    "ILP modeling · online scheduling\nvLLM integration · evaluation",
    0.66,
    3.98,
    3.05,
    0.68,
    size=13,
    color=LIGHT_BLUE,
)
add_text(
    slide, "WU CAO · 11036000", 0.66, 5.72, 3.0, 0.32, size=15, color=WHITE, bold=True
)
add_text(
    slide,
    "MSc High Performance Computing Engineering\nA.Y. 2026–2027",
    0.66,
    6.14,
    3.15,
    0.68,
    size=11,
    color=LIGHT_BLUE,
)
add_image_fit(slide, LOGO, 4.9, 0.58, 3.55, 1.5)
add_box(slide, 5.0, 2.4, 7.55, 2.45, fill=BG, line=LIGHT_BLUE)
add_text(
    slide,
    "Unbatched",
    5.32,
    2.66,
    1.1,
    0.25,
    size=10,
    color=RED,
    bold=True,
    align=PP_ALIGN.CENTER,
)
for i, label in enumerate(["R0", "R1", "R2", "R3"]):
    node(
        slide,
        label,
        5.38 + (i % 2) * 1.08,
        3.03 + (i // 2) * 0.84,
        color=(BLUE, CYAN, PURPLE, ORANGE)[i],
        size=0.48,
    )
for y, color in [(3.22, RED), (3.56, ORANGE), (3.9, PURPLE)]:
    add_line(slide, 5.88, y, 7.08, y, color=color, width=3, arrow=True)
add_text(
    slide,
    "all flows",
    6.12,
    4.18,
    0.75,
    0.2,
    size=8,
    color=RED,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "→",
    7.42,
    3.39,
    0.55,
    0.45,
    size=26,
    color=MUTED,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "Rank-disjoint batches",
    8.03,
    2.66,
    2.35,
    0.25,
    size=10,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER,
)
for i, (label, flows, color) in enumerate(
    [
        ("Batch 0", "R0→R1    R2→R3", BLUE),
        ("Batch 1", "R1→R2    R0→R3", PURPLE),
    ]
):
    y = 3.1 + i * 0.75
    add_box(slide, 8.03, y, 3.72, 0.55, fill=WHITE, line=color)
    add_text(slide, label, 8.18, y + 0.1, 0.72, 0.25, size=9, color=color, bold=True)
    add_text(
        slide,
        flows,
        9.0,
        y + 0.08,
        2.52,
        0.28,
        size=11,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
add_text(slide, "Advisor", 5.02, 5.45, 1.1, 0.25, size=10, color=BLUE, bold=True)
add_text(
    slide,
    "Prof. Massimo Tornatore",
    6.2,
    5.41,
    2.9,
    0.32,
    size=13,
    color=INK,
    bold=True,
)
add_text(slide, "Co-advisor", 5.02, 5.9, 1.1, 0.25, size=10, color=BLUE, bold=True)
add_text(
    slide,
    "Assistant Prof. Qiaolun Zhang",
    6.2,
    5.86,
    3.4,
    0.32,
    size=13,
    color=INK,
    bold=True,
)
add_text(
    slide,
    "Merged into vLLM · PR #52641",
    9.32,
    6.62,
    3.05,
    0.28,
    size=11,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.RIGHT,
    url=PR_URL,
)


# 02 — roadmap and research questions
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(slide, "Roadmap and research questions", 2, "Overview")
questions = [
    ("RQ1", "How should migration contention be represented?", CYAN),
    ("RQ2", "Can a lightweight scheduler enforce the model online?", BLUE),
    ("RQ3", "Can it preserve vLLM migration semantics?", PURPLE),
    ("RQ4", "Does it improve real serving at acceptable CPU cost?", GREEN),
]
for i, (tag, question, color) in enumerate(questions):
    x = 0.7 + (i % 2) * 6.18
    y = 1.72 + (i // 2) * 1.38
    add_box(slide, x, y, 5.72, 1.02, fill=WHITE, line=color, width=1.5)
    add_text(
        slide,
        tag,
        x + 0.18,
        y + 0.18,
        0.72,
        0.32,
        size=13,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide, question, x + 1.02, y + 0.12, 4.45, 0.48, size=15, color=INK, bold=True
    )
section_marker(slide, ["Background", "Model", "Algorithm", "Evaluation"], 0)
add_text(slide, "Narrative", 0.72, 4.72, 1.1, 0.26, size=10, color=MUTED, bold=True)
stages = [
    ("Observe", "migration bursts", RED),
    ("Model", "link contention", CYAN),
    ("Design", "rank-disjoint batches", BLUE),
    ("Integrate", "async vLLM EPLB", PURPLE),
    ("Validate", "E2E + NIC + trace", GREEN),
]
for i, (verb, noun, color) in enumerate(stages):
    x = 0.78 + i * 2.47
    add_box(slide, x, 5.12, 2.06, 0.85, fill=WHITE, line=color)
    add_text(
        slide,
        verb,
        x + 0.1,
        5.23,
        1.86,
        0.26,
        size=12,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        noun,
        x + 0.1,
        5.52,
        1.86,
        0.22,
        size=9,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    if i < len(stages) - 1:
        add_line(
            slide, x + 2.07, 5.54, x + 2.39, 5.54, color=MUTED, width=1.5, arrow=True
        )


# 03 — MoE and expert parallelism
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "MoE scales capacity through sparse expert activation",
    3,
    "Background",
    "Expert parallelism distributes the expert weights across ranks",
)
add_box(slide, 0.62, 1.72, 4.0, 4.92, fill=WHITE, line=BLUE)
add_text(slide, "TOKEN ROUTING", 0.92, 1.94, 1.8, 0.25, size=10, color=BLUE, bold=True)
for i, token in enumerate(["t₁", "t₂", "t₃", "t₄"]):
    node(slide, token, 0.95, 2.5 + i * 0.75, color=CYAN, size=0.52)
add_box(slide, 2.1, 3.1, 1.45, 1.08, fill=LIGHT_BLUE, line=BLUE)
add_text(
    slide,
    "Router\nTop-k",
    2.15,
    3.25,
    1.35,
    0.62,
    size=16,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
for i in range(4):
    add_line(
        slide,
        1.47,
        2.76 + i * 0.75,
        2.1,
        3.42 + (i % 2) * 0.24,
        color=GRID,
        width=1.2,
        arrow=True,
    )
for i, (expert, color) in enumerate([("E₀", BLUE), ("E₁", GREEN), ("E₂", ORANGE)]):
    node(slide, expert, 3.72, 2.45 + i * 1.16, color=color, size=0.64)
    add_line(
        slide, 3.55, 3.48, 3.72, 2.77 + i * 1.16, color=color, width=1.2, arrow=True
    )
add_text(
    slide,
    "Only a subset of experts runs for each token",
    0.92,
    5.72,
    3.35,
    0.46,
    size=13,
    color=INK,
    bold=True,
    align=PP_ALIGN.CENTER,
)

add_box(slide, 4.92, 1.72, 7.8, 4.92, fill=WHITE, line=GREEN)
add_text(
    slide,
    "EXPERT PARALLEL PLACEMENT",
    5.22,
    1.94,
    2.7,
    0.25,
    size=10,
    color=GREEN,
    bold=True,
)
for rank in range(4):
    x = 5.25 + rank * 1.8
    add_box(slide, x, 2.44, 1.43, 2.52, fill=BG, line=GRID)
    add_text(
        slide,
        f"Rank {rank}",
        x + 0.08,
        2.57,
        1.27,
        0.26,
        size=11,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    for slot in range(3):
        color = (BLUE, GREEN, ORANGE, PURPLE)[(rank + slot) % 4]
        add_box(
            slide,
            x + 0.24,
            3.02 + slot * 0.56,
            0.95,
            0.4,
            fill=color,
            line=color,
            radius=False,
        )
        add_text(
            slide,
            f"E{(rank * 3 + slot) % 8}",
            x + 0.24,
            3.07 + slot * 0.56,
            0.95,
            0.24,
            size=10,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
add_text(
    slide,
    "Routing is input-dependent",
    5.35,
    5.35,
    2.7,
    0.38,
    size=14,
    color=RED,
    bold=True,
)
add_text(
    slide,
    "→ expert popularity changes over time\n→ some ranks become overloaded",
    8.0,
    5.2,
    4.25,
    0.68,
    size=14,
    color=INK,
)
add_source(slide, "Background: Shazeer et al.; Switch Transformer; vLLM EPLB docs")


# 04 — EPLB and migration interference
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide, "Async EPLB creates network interference during migration", 4, "Background"
)
add_box(slide, 0.62, 1.7, 12.08, 2.18, fill=WHITE, line=BLUE)
steps = [
    ("1", "Collect\nexpert load", BLUE),
    ("2", "Compute new\nplacement", CYAN),
    ("3", "Migrate\nexpert weights", ORANGE),
    ("4", "Activate at\nsafe boundary", GREEN),
]
for i, (number, label, color) in enumerate(steps):
    x = 0.95 + i * 3.0
    node(slide, number, x, 2.18, color=color, size=0.55)
    add_text(slide, label, x + 0.72, 2.06, 1.8, 0.72, size=15, color=INK, bold=True)
    if i < len(steps) - 1:
        add_line(
            slide, x + 2.45, 2.47, x + 2.83, 2.47, color=MUTED, width=2, arrow=True
        )
add_text(
    slide,
    "Foreground inference continues across the entire async cycle",
    2.88,
    3.18,
    7.55,
    0.32,
    size=13,
    color=BLUE,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_line(slide, 1.25, 3.58, 11.95, 3.58, color=BLUE, width=4, arrow=True)

add_box(slide, 0.62, 4.15, 5.68, 2.23, fill=LIGHT_RED, line=RED)
add_text(
    slide, "UNBATCHED MIGRATION", 0.92, 4.39, 2.5, 0.26, size=11, color=RED, bold=True
)
for pair, (x1, y1, x2, y2) in enumerate(
    [
        (1.18, 4.98, 5.55, 4.98),
        (1.18, 5.28, 5.55, 5.28),
        (1.18, 5.58, 5.55, 5.58),
    ]
):
    add_line(
        slide, x1, y1, x2, y2, color=(RED, ORANGE, PURPLE)[pair], width=5, arrow=True
    )
add_text(
    slide,
    "many transfers injected together → NIC burst",
    1.3,
    5.84,
    4.1,
    0.34,
    size=13,
    color=INK,
    bold=True,
    align=PP_ALIGN.CENTER,
)

add_box(slide, 6.58, 4.15, 6.12, 2.23, fill=LIGHT_GREEN, line=GREEN)
add_text(
    slide, "RESEARCH QUESTION", 6.9, 4.39, 2.2, 0.26, size=11, color=GREEN, bold=True
)
add_text(
    slide,
    "Can migration traffic be reshaped\nwithout changing the placement?",
    7.15,
    4.92,
    5.0,
    0.8,
    size=21,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "Target: protect inference—not merely finish migration first",
    7.05,
    5.89,
    5.2,
    0.3,
    size=12,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# 05 — objective distinction
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "The correct objective depends on the execution mode",
    5,
    "Problem motivation",
    "Migration completion time and serving quality are different objectives",
)
add_box(slide, 0.72, 1.75, 5.8, 4.72, fill=WHITE, line=ORANGE, width=1.5)
add_text(
    slide, "SYNCHRONOUS EPLB", 1.05, 2.02, 2.6, 0.3, size=12, color=ORANGE, bold=True
)
add_text(slide, "Inference", 1.05, 2.68, 1.25, 0.28, size=11, color=MUTED, bold=True)
add_box(slide, 2.38, 2.62, 1.2, 0.44, fill=BLUE, line=BLUE, radius=False)
add_box(slide, 3.58, 2.62, 1.62, 0.44, fill=LIGHT_RED, line=RED, radius=False)
add_box(slide, 5.2, 2.62, 0.75, 0.44, fill=BLUE, line=BLUE, radius=False)
add_text(
    slide,
    "paused",
    3.7,
    2.67,
    1.35,
    0.24,
    size=9,
    color=RED,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(slide, "Migration", 1.05, 3.42, 1.25, 0.28, size=11, color=MUTED, bold=True)
add_box(slide, 3.58, 3.37, 1.62, 0.44, fill=ORANGE, line=ORANGE, radius=False)
add_text(slide, "Objective", 1.05, 4.28, 1.25, 0.28, size=11, color=ORANGE, bold=True)
add_text(
    slide,
    "Finish migration as fast as possible",
    2.38,
    4.1,
    3.45,
    0.68,
    size=17,
    color=INK,
    bold=True,
)
add_text(
    slide,
    "Batching is disabled in the implementation",
    1.15,
    5.65,
    4.95,
    0.34,
    size=12,
    color=ORANGE,
    bold=True,
    align=PP_ALIGN.CENTER,
)

add_box(slide, 6.82, 1.75, 5.8, 4.72, fill=WHITE, line=GREEN, width=1.5)
add_text(
    slide, "ASYNCHRONOUS EPLB", 7.15, 2.02, 2.7, 0.3, size=12, color=GREEN, bold=True
)
add_text(slide, "Inference", 7.15, 2.68, 1.25, 0.28, size=11, color=MUTED, bold=True)
add_box(slide, 8.48, 2.62, 3.58, 0.44, fill=BLUE, line=BLUE, radius=False)
add_text(
    slide,
    "continues",
    9.53,
    2.67,
    1.4,
    0.24,
    size=9,
    color=WHITE,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(slide, "Migration", 7.15, 3.42, 1.25, 0.28, size=11, color=MUTED, bold=True)
for i in range(3):
    add_box(
        slide, 8.48 + i * 1.2, 3.37, 0.91, 0.44, fill=GREEN, line=GREEN, radius=False
    )
add_text(slide, "Objective", 7.15, 4.28, 1.25, 0.28, size=11, color=GREEN, bold=True)
add_text(
    slide,
    "Reduce interference with active requests",
    8.48,
    4.1,
    3.45,
    0.68,
    size=17,
    color=INK,
    bold=True,
)
add_text(
    slide,
    "Opt-in batching trades migration speed for serving headroom",
    7.25,
    5.65,
    4.95,
    0.34,
    size=12,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# 06 — formal problem statement
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "Problem statement: schedule flows, not individual tensors",
    6,
    "System model",
)
cards = [
    (
        "GIVEN",
        "Old/new expert placements\n"
        "Network links and capacities\n"
        "Existing inference load",
        BLUE,
    ),
    (
        "DERIVE",
        "Remote instruction (s, d, e)\n"
        "Coalesced flow p = (s, d, Eₚ)\n"
        "Flow size wₚ = |Eₚ|",
        CYAN,
    ),
    ("DECIDE", "Assign each directed flow\nto one migration batch", PURPLE),
    (
        "OPTIMIZE",
        "Minimize worst modeled\ninference time under\nlink contention",
        GREEN,
    ),
]
for i, (tag, body, color) in enumerate(cards):
    x = 0.63 + i * 3.12
    add_box(slide, x, 1.82, 2.82, 2.55, fill=WHITE, line=color, width=1.5)
    add_text(
        slide,
        tag,
        x + 0.18,
        2.02,
        2.46,
        0.28,
        size=11,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        body,
        x + 0.22,
        2.53,
        2.38,
        1.35,
        size=14,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    if i < len(cards) - 1:
        add_line(
            slide, x + 2.84, 3.08, x + 3.06, 3.08, color=MUTED, width=1.5, arrow=True
        )

add_box(slide, 0.82, 4.78, 11.68, 1.23, fill=LIGHT_BLUE, line=BLUE)
add_text(
    slide, "Hard constraint", 1.08, 4.98, 1.65, 0.3, size=12, color=BLUE, bold=True
)
add_text(
    slide,
    "Each rank participates in at most one coalesced directed flow per batch",
    2.82,
    4.9,
    9.23,
    0.48,
    size=18,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "Independent rank pairs remain parallel; "
    "conflicting endpoint traffic is serialized",
    2.82,
    5.4,
    9.23,
    0.3,
    size=12,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)
section_marker(slide, ["Background", "Model", "Algorithm", "Evaluation"], 1)


# 07 — ILP formulation
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "ILP formulation exposes the contention trade-off",
    7,
    "Offline model",
    "The ILP is explanatory; it is not executed in the serving path",
)
add_box(slide, 0.62, 1.68, 3.18, 4.87, fill=WHITE, line=BLUE)
add_text(
    slide, "DECISION VARIABLES", 0.9, 1.92, 2.6, 0.26, size=10, color=BLUE, bold=True
)
add_text(
    slide,
    "xₚ,ᵦ",
    0.95,
    2.43,
    0.85,
    0.4,
    size=24,
    color=NAVY,
    bold=True,
    font="Cambria Math",
)
add_text(
    slide, "flow p assigned to batch b", 1.84, 2.42, 1.65, 0.44, size=12, color=INK
)
add_text(
    slide,
    "yᵦ",
    0.95,
    3.21,
    0.85,
    0.4,
    size=24,
    color=NAVY,
    bold=True,
    font="Cambria Math",
)
add_text(slide, "batch b is active", 1.84, 3.2, 1.65, 0.44, size=12, color=INK)
add_text(
    slide,
    "Qₗ,ᵦ",
    0.95,
    3.99,
    0.85,
    0.4,
    size=24,
    color=NAVY,
    bold=True,
    font="Cambria Math",
)
add_text(slide, "migration load on link l", 1.84, 3.98, 1.65, 0.44, size=12, color=INK)
add_box(slide, 0.9, 4.82, 2.62, 1.05, fill=LIGHT_ORANGE, line=ORANGE)
add_text(
    slide,
    "Tie-break only",
    1.05,
    4.95,
    2.3,
    0.25,
    size=10,
    color=ORANGE,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "After fixing the optimum,\nminimize active batches",
    1.05,
    5.24,
    2.3,
    0.48,
    size=12,
    color=INK,
    bold=True,
    align=PP_ALIGN.CENTER,
)

formula_panels = [
    ("1 · Every flow executes exactly once", "assign", 1.8, BLUE),
    ("2 · Rank-disjoint endpoint constraint", "endpoint", 3.08, PURPLE),
    ("3 · Worst-link serving objective", "objective", 4.36, GREEN),
]
for label, key, y, color in formula_panels:
    add_box(slide, 4.08, y, 8.5, 1.02, fill=WHITE, line=color)
    add_text(slide, label, 4.33, y + 0.1, 3.1, 0.24, size=10, color=color, bold=True)
    add_equation(
        slide, key, 6.02, y + 0.16, 6.2, 0.68, size=16 if key != "objective" else 14
    )
add_text(
    slide,
    "Primary objective: lower the maximum link pressure seen by concurrent inference",
    4.22,
    5.7,
    8.15,
    0.36,
    size=14,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_source(slide, "Offline model: SciPy MILP / HiGHS")


# 08 — ILP example and result
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "The constrained optimum reduces contention, not migration time",
    8,
    "Offline model",
    "Illustrative 8-rank hotspot with normalized capacities",
)
add_box(slide, 0.55, 1.65, 4.15, 3.95, fill=WHITE, line=RED)
add_text(slide, "UNBATCHED", 0.8, 1.84, 1.2, 0.25, size=10, color=RED, bold=True)
add_image_fit(slide, ILP / "example_hotspot_before.png", 0.68, 2.04, 3.89, 3.38)
add_box(slide, 4.92, 1.65, 4.15, 3.95, fill=WHITE, line=GREEN)
add_text(
    slide,
    "ILP SCHEDULE · 4 BATCHES",
    5.17,
    1.84,
    2.35,
    0.25,
    size=10,
    color=GREEN,
    bold=True,
)
add_image_fit(slide, ILP / "example_hotspot_after.png", 5.05, 2.04, 3.89, 3.38)

rows = [
    ["Metric", "Off", "ILP"],
    ["Modeled inference time", "4.75", "2.25"],
    ["Peak link utilization", "1.50", "0.50"],
    ["Migration time", "1.55", "2.20"],
    ["Batches", "1", "4"],
]
add_table(slide, rows, [2.2, 0.72, 0.72], 9.33, 1.82, 0.66, font_size=10)
add_box(slide, 9.33, 5.34, 3.64, 0.92, fill=LIGHT_BLUE, line=BLUE)
add_text(slide, "Key insight", 9.52, 5.45, 0.9, 0.22, size=9, color=BLUE, bold=True)
add_text(
    slide,
    "More batches can make serving faster\nwhile migration itself finishes later",
    10.25,
    5.34,
    2.5,
    0.72,
    size=11,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_source(slide, "Source: ILP/example_hotspot.json")


# 09 — production requirements
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "From an exact model to a production scheduler",
    9,
    "Online design",
    "An online MILP would optimize an approximation "
    "while adding dependencies and latency",
)
requirements = [
    (
        "1",
        "Semantic preservation",
        "Same placement and replica source assignment",
        BLUE,
    ),
    ("2", "Conflict freedom", "No rank in two directed flows within a batch", PURPLE),
    ("3", "Determinism", "All ranks derive the same schedule locally", CYAN),
    ("4", "Low CPU cost", "No optimizer in a periodic serving path", GREEN),
    (
        "5",
        "NIXL compatibility",
        "Batch boundaries survive communicator execution",
        ORANGE,
    ),
    ("6", "Safe rollout", "Opt-in and async-only; legacy path retained", RED),
]
for i, (number, title, body, color) in enumerate(requirements):
    x = 0.68 + (i % 3) * 4.16
    y = 1.78 + (i // 3) * 1.78
    add_box(slide, x, y, 3.78, 1.38, fill=WHITE, line=color)
    node(slide, number, x + 0.18, y + 0.23, color=color, size=0.48)
    add_text(
        slide, title, x + 0.82, y + 0.18, 2.65, 0.32, size=14, color=NAVY, bold=True
    )
    add_text(slide, body, x + 0.82, y + 0.58, 2.65, 0.48, size=11, color=MUTED)
add_box(slide, 1.25, 5.55, 10.75, 0.72, fill=LIGHT_GREEN, line=GREEN)
add_text(
    slide,
    "Chosen abstraction: deterministic first-fit edge coloring "
    "of coalesced directed rank flows",
    1.5,
    5.68,
    10.25,
    0.36,
    size=17,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
section_marker(slide, ["Background", "Model", "Algorithm", "Evaluation"], 2)


# 10 — algorithm pipeline
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "Online scheduler: five transformations from placement to batches",
    10,
    "Algorithm",
    "The implementation preserves placement semantics and changes only transfer order",
)
pipeline = [
    ("1", "Compare placements", "old/new local\nexpert sets", BLUE),
    ("2", "Find remote gains", "exclude experts\nalready local", CYAN),
    ("3", "Choose sources", "balance destinations\nover replicas", PURPLE),
    ("4", "Coalesce flows", "group experts by\ndirected (src,dst)", ORANGE),
    ("5", "First-fit grouping", "build rank-disjoint\nmigration batches", GREEN),
]
for i, (number, title, detail, color) in enumerate(pipeline):
    x = 0.58 + i * 2.52
    add_box(slide, x, 1.82, 2.18, 2.27, fill=WHITE, line=color, width=1.5)
    node(slide, number, x + 0.16, 2.0, color=color, size=0.48)
    add_text(slide, title, x + 0.72, 1.96, 1.28, 0.4, size=13, color=NAVY, bold=True)
    add_text(
        slide,
        detail,
        x + 0.18,
        2.68,
        1.82,
        0.82,
        size=13,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    if i < len(pipeline) - 1:
        add_line(
            slide, x + 2.2, 2.96, x + 2.46, 2.96, color=MUTED, width=1.5, arrow=True
        )

add_box(slide, 0.82, 4.47, 5.75, 1.58, fill=LIGHT_BLUE, line=BLUE)
add_text(
    slide,
    "Example placement difference",
    1.05,
    4.65,
    2.35,
    0.28,
    size=11,
    color=BLUE,
    bold=True,
)
add_text(
    slide,
    "Rank 1: old {E₀,E₂}  →  new {E₀,E₃}",
    1.12,
    5.08,
    4.9,
    0.32,
    size=16,
    color=NAVY,
    bold=True,
)
add_text(
    slide,
    "E₀ is local; only E₃ becomes a remote migration",
    1.12,
    5.49,
    4.9,
    0.3,
    size=12,
    color=MUTED,
)
add_box(slide, 6.82, 4.47, 5.68, 1.58, fill=LIGHT_ORANGE, line=ORANGE)
add_text(
    slide, "Why coalesce?", 7.05, 4.65, 1.45, 0.28, size=11, color=ORANGE, bold=True
)
for x, experts in [(7.15, "0→1 · E₃"), (8.64, "0→1 · E₇")]:
    add_box(slide, x, 5.06, 1.3, 0.4, fill=WHITE, line=ORANGE)
    add_text(
        slide,
        experts,
        x + 0.05,
        5.12,
        1.2,
        0.24,
        size=11,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
add_text(
    slide,
    "→",
    10.02,
    5.05,
    0.35,
    0.34,
    size=18,
    color=ORANGE,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_box(slide, 10.42, 5.02, 1.62, 0.46, fill=WHITE, line=GREEN)
add_text(
    slide,
    "0→1 · {E₃,E₇}",
    10.48,
    5.1,
    1.5,
    0.25,
    size=11,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "One contention decision per directed rank pair",
    7.08,
    5.49,
    5.05,
    0.3,
    size=12,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)
add_source(slide, "Heuristic inspired by OmniInfer placement grouping", OMNI_URL)


# 11 — worked greedy example
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "Greedy first-fit groups independent flows in the same batch",
    11,
    "Algorithm",
    "A batch is a matching in the rank communication graph",
)
add_box(slide, 0.62, 1.72, 4.0, 4.9, fill=WHITE, line=BLUE)
add_text(slide, "INPUT ORDER", 0.9, 1.93, 1.35, 0.24, size=10, color=BLUE, bold=True)
flows = [
    ("F₀", "0 → 1", "{E₁,E₇}", BLUE),
    ("F₁", "2 → 3", "{E₃}", GREEN),
    ("F₂", "1 → 2", "{E₄}", PURPLE),
    ("F₃", "0 → 3", "{E₉}", ORANGE),
]
for i, (name, pair, experts, color) in enumerate(flows):
    y = 2.38 + i * 0.8
    add_box(slide, 0.94, y, 3.34, 0.59, fill=BG, line=color)
    add_text(slide, name, 1.08, y + 0.11, 0.55, 0.26, size=12, color=color, bold=True)
    add_text(
        slide,
        pair,
        1.75,
        y + 0.09,
        0.85,
        0.28,
        size=15,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        experts,
        2.72,
        y + 0.1,
        1.28,
        0.27,
        size=12,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
add_equation(slide, "mask", 0.88, 5.72, 3.5, 0.58, size=11)

add_box(slide, 4.92, 1.72, 7.78, 4.9, fill=WHITE, line=GREEN)
add_text(
    slide, "FIRST-FIT DECISIONS", 5.22, 1.93, 2.2, 0.24, size=10, color=GREEN, bold=True
)
decisions = [
    ("F₀", "Batch 0", "uses ranks {0,1}", BLUE),
    ("F₁", "Batch 0", "disjoint ranks {2,3}", GREEN),
    ("F₂", "Batch 1", "conflicts with F₀ and F₁", PURPLE),
    ("F₃", "Batch 1", "disjoint from F₂", ORANGE),
]
for i, (flow, batch, reason, color) in enumerate(decisions):
    y = 2.36 + i * 0.61
    node(slide, flow, 5.22, y, color=color, size=0.48)
    add_text(slide, batch, 5.88, y + 0.05, 1.08, 0.28, size=12, color=NAVY, bold=True)
    add_text(slide, reason, 7.02, y + 0.03, 2.8, 0.3, size=11, color=MUTED)

for batch_index, (y, batch_color, entries) in enumerate(
    [
        (5.08, BLUE, [("0", "1"), ("2", "3")]),
        (5.72, PURPLE, [("1", "2"), ("0", "3")]),
    ]
):
    add_text(
        slide,
        f"Batch {batch_index}",
        9.72,
        y,
        0.8,
        0.24,
        size=10,
        color=batch_color,
        bold=True,
    )
    for j, (src, dst) in enumerate(entries):
        x = 10.56 + j * 0.9
        node(slide, src, x, y - 0.05, color=batch_color, size=0.38)
        add_line(
            slide,
            x + 0.4,
            y + 0.14,
            x + 0.63,
            y + 0.14,
            color=batch_color,
            width=1.5,
            arrow=True,
        )
        node(slide, dst, x + 0.65, y - 0.05, color=batch_color, size=0.38)
add_text(
    slide,
    "Reciprocal flows are distinct directed flows and cannot share a batch",
    5.24,
    6.2,
    6.98,
    0.28,
    size=11,
    color=RED,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# 12 — correctness, complexity, tests
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "Correctness and cost are explicit—not empirical assumptions",
    12,
    "Algorithm analysis",
)
invariants = [
    ("Completeness", "Every remote gain appears\nin exactly one flow and batch", BLUE),
    (
        "Conflict freedom",
        "Endpoint masks prohibit\nshared ranks inside a batch",
        PURPLE,
    ),
    (
        "Semantic preservation",
        "Local copies and balanced\nreplica sources are unchanged",
        GREEN,
    ),
]
for i, (title, body, color) in enumerate(invariants):
    x = 0.66 + i * 4.15
    add_box(slide, x, 1.72, 3.78, 1.28, fill=WHITE, line=color)
    add_text(
        slide,
        title,
        x + 0.18,
        1.89,
        3.42,
        0.3,
        size=14,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        body,
        x + 0.26,
        2.3,
        3.26,
        0.46,
        size=10,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

add_box(slide, 0.66, 3.32, 5.6, 2.6, fill=WHITE, line=BLUE)
add_text(slide, "COMPLEXITY", 0.94, 3.52, 1.2, 0.25, size=10, color=BLUE, bold=True)
add_equation(slide, "complexity", 1.0, 3.93, 4.9, 0.64, size=16)
add_text(
    slide,
    "S: slots · M: remote migrations · F: directed flows · K: batches",
    1.03,
    4.72,
    4.85,
    0.28,
    size=11,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "For fixed small EP degree, practical cost is close to linear.",
    1.03,
    5.18,
    4.85,
    0.34,
    size=13,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)

add_box(slide, 6.55, 3.32, 6.15, 2.6, fill=WHITE, line=GREEN)
add_text(
    slide, "12 UNIT TESTS", 6.83, 3.52, 1.45, 0.25, size=10, color=GREEN, bold=True
)
tests = [
    "deterministic grouping",
    "directed-pair coalescing",
    "random completeness",
    "local-copy exclusion",
    "replica balancing",
    "multi-layer equivalence",
    "mock communicator execution",
    "sync/async flag boundaries",
]
for i, test in enumerate(tests):
    x = 6.9 + (i % 2) * 2.75
    y = 3.93 + (i // 2) * 0.42
    add_text(slide, "✓", x, y, 0.25, 0.22, size=11, color=GREEN, bold=True)
    add_text(slide, test, x + 0.28, y - 0.01, 2.35, 0.26, size=10, color=INK)
add_text(
    slide,
    "Public behavior is tested; the ILP is not a runtime dependency.",
    6.9,
    5.61,
    5.5,
    0.26,
    size=11,
    color=MUTED,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# 13 — vLLM integration
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "vLLM separates cycle-wide scheduling from layer-wise transfer",
    13,
    "System implementation",
)
add_box(slide, 0.6, 1.66, 4.0, 4.98, fill=LIGHT_BLUE, line=BLUE)
add_text(
    slide,
    "ONCE PER REBALANCE CYCLE",
    0.9,
    1.92,
    2.45,
    0.26,
    size=10,
    color=BLUE,
    bold=True,
)
add_box(slide, 1.02, 2.52, 3.18, 0.75, fill=WHITE, line=BLUE)
add_text(
    slide,
    "old/new placement arrays\nfor all 48 MoE layers",
    1.2,
    2.62,
    2.82,
    0.5,
    size=14,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_line(slide, 2.61, 3.28, 2.61, 3.65, color=BLUE, width=2, arrow=True)
add_box(slide, 1.02, 3.68, 3.18, 0.88, fill=BLUE, line=BLUE)
add_text(
    slide,
    "schedule_migration_batches_\nfor_layers(...)  × 1 call",
    1.17,
    3.8,
    2.88,
    0.56,
    size=14,
    color=WHITE,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_line(slide, 2.61, 4.58, 2.61, 4.95, color=BLUE, width=2, arrow=True)
add_text(
    slide,
    "48 independent schedules",
    1.47,
    4.98,
    2.3,
    0.46,
    size=14,
    color=BLUE,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "CPU work is aggregated; expert payloads are not.",
    1.0,
    5.62,
    3.2,
    0.46,
    size=12,
    color=INK,
    bold=True,
    align=PP_ALIGN.CENTER,
)

add_box(slide, 4.9, 1.66, 7.82, 4.98, fill=WHITE, line=GREEN)
add_text(
    slide,
    "LAYER-WISE NIXL EXECUTION",
    5.2,
    1.92,
    2.5,
    0.26,
    size=10,
    color=GREEN,
    bold=True,
)
for layer in range(4):
    y = 2.42 + layer * 0.88
    add_text(
        slide,
        f"Layer {layer if layer < 3 else '47'}",
        5.18,
        y + 0.1,
        0.72,
        0.25,
        size=10,
        color=NAVY,
        bold=True,
    )
    stages = [
        ("set context", BLUE, 1.22),
        ("add send/recv", ORANGE, 1.5),
        ("execute batch", GREEN, 1.36),
        ("activate", PURPLE, 1.05),
    ]
    x = 5.98
    for j, (label, color, width) in enumerate(stages):
        add_box(slide, x, y, width, 0.49, fill=color, line=color, radius=False)
        add_text(
            slide,
            label,
            x + 0.04,
            y + 0.08,
            width - 0.08,
            0.25,
            size=9,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        if j < len(stages) - 1:
            add_line(
                slide,
                x + width,
                y + 0.25,
                x + width + 0.16,
                y + 0.25,
                color=MUTED,
                width=1,
                arrow=True,
            )
        x += width + 0.18
add_box(slide, 5.2, 5.96, 7.12, 0.43, fill=LIGHT_ORANGE, line=ORANGE)
add_text(
    slide,
    "NIXL execute() drains the batch and clears context → restore context per batch",
    5.34,
    6.04,
    6.84,
    0.24,
    size=10,
    color=INK,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# 14 — evaluation setup
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "Controlled A/B evaluation on a four-node cluster",
    14,
    "Evaluation methodology",
    "Fresh server per condition; batching is the only changed switch",
)
add_box(slide, 0.62, 1.7, 7.15, 4.87, fill=WHITE, line=BLUE)
for i in range(4):
    x = 0.97 + i * 1.65
    add_box(slide, x, 2.15, 1.27, 1.15, fill=LIGHT_BLUE, line=BLUE)
    add_text(
        slide,
        f"Node {i}",
        x + 0.08,
        2.27,
        1.11,
        0.24,
        size=10,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "RTX 6000\n24 GB",
        x + 0.08,
        2.57,
        1.11,
        0.48,
        size=12,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_line(slide, x + 0.64, 3.31, x + 0.64, 3.87, color=GREEN, width=3)
add_line(slide, 1.6, 3.87, 6.55, 3.87, color=GREEN, width=4)
add_text(
    slide,
    "10 GbE · NIXL 1.3.2 over UCX/TCP · no RDMA",
    1.25,
    4.05,
    5.75,
    0.32,
    size=13,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "Qwen3-30B-A3B-Instruct-2507 · FP16 · TP=EP=4",
    1.18,
    4.62,
    5.9,
    0.32,
    size=13,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "Async EPLB · window 50 · step interval 50 · 16 redundant experts",
    1.1,
    5.1,
    6.05,
    0.3,
    size=11,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)
add_box(slide, 1.12, 5.58, 6.0, 0.54, fill=LIGHT_GREEN, line=GREEN)
add_text(
    slide,
    "Prefix cache enabled · eager execution · 500/500 requests completed",
    1.24,
    5.68,
    5.75,
    0.26,
    size=11,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER,
)

add_box(slide, 8.05, 1.7, 4.65, 4.87, fill=WHITE, line=PURPLE)
add_text(slide, "WORKLOAD", 8.35, 1.94, 1.2, 0.25, size=10, color=PURPLE, bold=True)
params = [
    ("Dataset", "random · seed 0"),
    ("Prompts", "500"),
    ("Concurrency", "32"),
    ("Prefix / input", "300 / 200 tokens"),
    ("Output", "300 tokens · ignore EOS"),
    ("Reported", "P50 and P99"),
]
for i, (label, value) in enumerate(params):
    y = 2.42 + i * 0.54
    add_text(slide, label, 8.37, y, 1.25, 0.25, size=10, color=MUTED, bold=True)
    add_text(slide, value, 9.66, y - 0.02, 2.62, 0.28, size=12, color=INK, bold=True)
add_box(slide, 8.35, 5.77, 4.05, 0.48, fill=BG, line=GRID)
add_text(
    slide,
    "OFF vs ON · same request stream",
    8.5,
    5.86,
    3.75,
    0.24,
    size=11,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
section_marker(slide, ["Background", "Model", "Algorithm", "Evaluation"], 3)
add_source(slide, "Raw data: results/serving_nixl_20260911_async_random500")


# 15 — complete E2E table
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "End-to-end results: every reported serving metric improves",
    15,
    "Evaluation results",
    "Positive improvement means higher throughput or lower latency",
)
rows = [
    ["Metric", "Batching off", "Batching on", "Improvement"],
    ["Output throughput", "31.54 tok/s", "32.26 tok/s", "+2.28%"],
    ["TTFT P50", "49,180.71 ms", "48,599.04 ms", "+1.18%"],
    ["TTFT P99", "75,740.25 ms", "74,756.17 ms", "+1.30%"],
    ["TPOT P50", "828.53 ms", "812.44 ms", "+1.94%"],
    ["TPOT P99", "1,007.82 ms", "999.01 ms", "+0.87%"],
    ["E2EL P50", "303,462.69 ms", "293,821.76 ms", "+3.18%"],
    ["E2EL P99", "322,184.88 ms", "317,351.61 ms", "+1.50%"],
]
add_table(slide, rows, [2.35, 2.55, 2.55, 1.75], 0.82, 1.72, 0.56, font_size=11)
add_box(slide, 10.35, 1.72, 2.28, 4.48, fill=LIGHT_GREEN, line=GREEN)
add_text(
    slide,
    "MAIN RESULT",
    10.63,
    2.0,
    1.72,
    0.26,
    size=10,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "+2.28%",
    10.55,
    2.56,
    1.88,
    0.52,
    size=28,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "output throughput",
    10.62,
    3.08,
    1.75,
    0.3,
    size=10,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "−3.18%",
    10.55,
    3.76,
    1.88,
    0.52,
    size=28,
    color=BLUE,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "E2EL P50",
    10.62,
    4.28,
    1.75,
    0.3,
    size=10,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "No regression",
    10.55,
    5.17,
    1.88,
    0.34,
    size=15,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "in the reported async metrics",
    10.58,
    5.53,
    1.82,
    0.42,
    size=9,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)
add_box(slide, 0.82, 6.37, 11.8, 0.42, fill=LIGHT_BLUE, line=BLUE)
add_text(
    slide,
    "Interpretation: throttling migration reduces shared-resource "
    "interference during active serving",
    1.08,
    6.45,
    11.28,
    0.23,
    size=11,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_source(slide, "Source: e2e_summary.csv")


# 16 — E2E visualization
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "The benefit is consistent across throughput and latency",
    16,
    "Evaluation results",
)
add_box(slide, 0.58, 1.68, 9.25, 4.98, fill=WHITE, line=GREEN)
add_image_fit(slide, E2E / "e2e_improvement.png", 0.75, 1.88, 8.92, 4.58)
add_metric(slide, "+2.28%", "throughput", 10.15, 1.9, 2.3, color=GREEN)
add_metric(
    slide,
    "+1.18 / +1.30%",
    "TTFT P50 / P99",
    10.15,
    3.12,
    2.3,
    color=BLUE,
    value_size=16,
)
add_metric(
    slide,
    "+1.94 / +0.87%",
    "TPOT P50 / P99",
    10.15,
    4.34,
    2.3,
    color=PURPLE,
    value_size=16,
)
add_metric(
    slide,
    "+3.18 / +1.50%",
    "E2EL P50 / P99",
    10.15,
    5.56,
    2.3,
    color=CYAN,
    value_size=16,
)
add_source(slide, "Source: e2e_improvement.png")


# 17 — NIC mechanism
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "Lower NIC tail traffic supports the proposed mechanism",
    17,
    "Network evidence",
    "One-second RX+TX samples on the head-node eno2 interface",
)
add_box(slide, 0.55, 1.62, 10.02, 5.14, fill=WHITE, line=CYAN)
add_image_fit(slide, E2E / "nic_timeseries.png", 0.7, 1.79, 9.72, 4.78)
add_metric(slide, "599.81", "MB/s P99 · off", 10.82, 1.9, 1.92, color=RED)
add_metric(slide, "538.40", "MB/s P99 · on", 10.82, 3.25, 1.92, color=GREEN)
add_metric(slide, "−10.24%", "NIC P99", 10.82, 4.6, 1.92, color=CYAN)
add_text(
    slide,
    "Tail bursts are lower, leaving more headroom for inference traffic.",
    10.84,
    5.85,
    1.88,
    0.62,
    size=10,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_source(slide, "Source: nic.tsv and nic_summary.csv")


# 18 — profiling method
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide,
    "Profiling measures scheduler cost in the serving context",
    18,
    "Scheduler profiling",
    "Clean E2E and profiling are separate runs to avoid contaminating the A/B result",
)
add_box(slide, 0.7, 1.75, 12.0, 1.4, fill=WHITE, line=PURPLE)
timeline = [
    ("Inference", 0.0, 10.6, BLUE),
    ("Async EPLB worker", 1.15, 1.65, PURPLE),
    ("Layer-wise NIXL transfers", 3.25, 4.9, ORANGE),
]
for label, start, width, color in timeline:
    y = 2.0 + timeline.index((label, start, width, color)) * 0.31
    add_text(slide, label, 0.93, y, 1.55, 0.2, size=8, color=MUTED, bold=True)
    add_box(
        slide,
        2.58 + start * 0.72,
        y - 0.01,
        width * 0.72,
        0.2,
        fill=color,
        line=color,
        radius=False,
    )
add_text(
    slide,
    "NVTX scheduler range",
    3.44,
    2.32,
    1.36,
    0.2,
    size=7,
    color=WHITE,
    bold=True,
    align=PP_ALIGN.CENTER,
)

add_box(slide, 0.7, 3.52, 5.72, 2.45, fill=LIGHT_BLUE, line=BLUE)
add_text(slide, "CAPTURE", 0.98, 3.75, 1.0, 0.25, size=10, color=BLUE, bold=True)
add_bullets(
    slide,
    [
        "NVTX-only Nsight Systems on all four Ray workers",
        "Same model, EPLB settings, workload lengths and 500 prompts",
        "CPU sampling and context-switch tracing disabled",
        "34 cycle-wide scheduler calls captured per rank",
    ],
    1.0,
    4.12,
    5.05,
    1.55,
    size=12,
    spacing=6,
)

add_box(slide, 6.72, 3.52, 5.98, 2.45, fill=LIGHT_ORANGE, line=ORANGE)
add_text(
    slide,
    "TEMPORARY INSTRUMENTATION",
    7.0,
    3.75,
    2.55,
    0.25,
    size=10,
    color=ORANGE,
    bold=True,
)
add_bullets(
    slide,
    [
        "An explicit CUDA NVTX range exposes the background thread",
        "The range was absent from the clean E2E A/B experiment",
        "It was removed from the final production patch",
        "Only scheduler elapsed time—not transfer time—is attributed",
    ],
    7.02,
    4.12,
    5.28,
    1.55,
    size=12,
    spacing=6,
)
add_source(slide, "Raw traces: ...async_random500_profile/trace/rank*.nsys-rep")


# 19 — profile results and ratio
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_header(
    slide, "Scheduler cost is 0.285% of serving time saved", 19, "Scheduler profiling"
)
profile_rows = [
    ["Rank", "Calls", "P50/call", "P99/call", "Cumulative"],
    ["0", "34", "7.819 ms", "17.536 ms", "302.081 ms"],
    ["1", "34", "7.090 ms", "12.357 ms", "253.152 ms"],
    ["2", "34", "7.035 ms", "16.104 ms", "255.239 ms"],
    ["3", "34", "7.005 ms", "8.103 ms", "241.190 ms"],
]
add_table(
    slide, profile_rows, [0.72, 0.78, 1.36, 1.36, 1.55], 0.7, 1.74, 0.57, font_size=10
)
add_text(
    slide,
    "One call builds schedules for all 48 MoE layers",
    0.95,
    4.88,
    5.1,
    0.35,
    size=14,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "Ranks schedule concurrently → use the largest cumulative rank cost",
    0.95,
    5.36,
    5.1,
    0.42,
    size=11,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)

add_box(slide, 6.86, 1.74, 5.78, 4.8, fill=WHITE, line=GREEN)
add_text(
    slide,
    "COST / SAVED COMPARISON",
    7.16,
    1.98,
    2.62,
    0.25,
    size=10,
    color=GREEN,
    bold=True,
)
add_equation(slide, "ratio", 7.12, 2.39, 5.25, 0.78, size=15)
add_metric(
    slide,
    "302.081 ms",
    "largest cumulative scheduler cost",
    7.22,
    3.48,
    2.32,
    color=ORANGE,
)
add_metric(
    slide, "105.932 s", "clean E2E serving time saved", 9.92, 3.48, 2.32, color=GREEN
)
add_box(slide, 7.22, 4.83, 5.02, 1.15, fill=LIGHT_GREEN, line=GREEN)
add_text(
    slide,
    "0.285%",
    7.48,
    4.96,
    1.65,
    0.45,
    size=27,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "scheduler cost / serving time saved",
    9.05,
    5.02,
    2.9,
    0.34,
    size=12,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "This is a conservative cost comparison, not a causal decomposition.",
    7.28,
    6.11,
    4.86,
    0.26,
    size=9,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)
add_source(slide, "Sources: profile_summary.csv + clean E2E off/on duration")


# 20 — contributions and conclusion
slide = prs.slides.add_slide(BLANK)
add_bg(slide, WHITE)
add_box(slide, 0, 0, 3.95, 7.5, fill=NAVY, line=NAVY, radius=False)
add_text(
    slide,
    "CONCLUSION",
    0.62,
    0.65,
    2.2,
    0.3,
    size=12,
    color=RGBColor(163, 203, 232),
    bold=True,
)
add_text(
    slide,
    "Schedule migration\nas part of serving—\nnot as an isolated\ncopy task.",
    0.62,
    1.25,
    2.95,
    3.2,
    size=25,
    color=WHITE,
    bold=True,
    valign=MSO_ANCHOR.TOP,
)
add_text(slide, "Questions?", 0.62, 5.65, 2.6, 0.52, size=28, color=WHITE, bold=True)
add_text(
    slide,
    "vLLM PR #52641",
    0.63,
    6.48,
    2.35,
    0.28,
    size=11,
    color=LIGHT_BLUE,
    bold=True,
    url=PR_URL,
)

add_text(
    slide, "Thesis contributions", 4.25, 0.58, 3.5, 0.48, size=25, color=NAVY, bold=True
)
contributions = [
    (
        "Model",
        "Separates foreground interference\nfrom migration completion time",
        CYAN,
    ),
    ("Algorithm", "Rank-disjoint flow batching\nwith explicit invariants", BLUE),
    (
        "System",
        "Async vLLM integration: cycle-wide\nscheduling, layer-wise transfer",
        PURPLE,
    ),
    ("Evidence", "Throughput +2.28% · E2EL P50 −3.18%\nNIC P99 −10.24%", GREEN),
]
for i, (title, body, color) in enumerate(contributions):
    y = 1.36 + i * 1.1
    add_box(slide, 4.25, y, 5.65, 0.87, fill=BG, line=color)
    add_text(slide, title, 4.45, y + 0.15, 1.1, 0.28, size=13, color=color, bold=True)
    add_text(slide, body, 5.58, y + 0.08, 4.05, 0.58, size=10, color=INK, bold=True)

add_box(slide, 10.22, 1.36, 2.58, 4.17, fill=LIGHT_ORANGE, line=ORANGE)
add_text(
    slide,
    "LIMITATIONS",
    10.52,
    1.62,
    1.98,
    0.25,
    size=10,
    color=ORANGE,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "• one four-node cluster\n\n"
    "• one model and NIXL backend\n\n"
    "• one A/B pair per condition\n\n"
    "• synthetic fixed-length workload",
    10.53,
    2.15,
    1.95,
    2.55,
    size=12,
    color=INK,
    bold=True,
    valign=MSO_ANCHOR.TOP,
)
add_box(slide, 4.25, 6.08, 8.55, 0.68, fill=LIGHT_GREEN, line=GREEN)
add_text(
    slide,
    "Result: a measured, opt-in optimization merged into upstream vLLM",
    4.52,
    6.22,
    8.0,
    0.34,
    size=17,
    color=NAVY,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "WU CAO · MSc Thesis Defense · 2026–2027",
    9.45,
    7.12,
    3.3,
    0.2,
    size=8,
    color=MUTED,
    align=PP_ALIGN.RIGHT,
)


prs.core_properties.title = (
    "Contention-Aware Expert Migration for MoE Serving — MSc Thesis Defense"
)
prs.core_properties.subject = (
    "ILP formulation, online scheduling, vLLM integration, and evaluation"
)
prs.core_properties.author = "WU CAO"
prs.core_properties.keywords = (
    "MoE, expert parallelism, EPLB, migration batching, vLLM, NIXL"
)
prs.save(OUT)
print(OUT)
