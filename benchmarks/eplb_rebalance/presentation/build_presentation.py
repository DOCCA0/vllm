from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
ILP = ROOT / "benchmarks/eplb_rebalance/ILP"
E2E = ROOT / (
    "benchmarks/eplb_rebalance/results/"
    "serving_nixl_20260911_async_random500"
)
PROFILE = ROOT / (
    "benchmarks/eplb_rebalance/results/"
    "serving_nixl_20260911_async_random500_profile"
)

OUT = HERE / "eplb_migration_batching_pr52641.pptx"
PR_URL = "https://github.com/vllm-project/vllm/pull/52641"

BG = RGBColor(247, 249, 252)
PANEL = RGBColor(255, 255, 255)
PANEL_2 = RGBColor(237, 243, 250)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(22, 34, 54)
MUTED = RGBColor(89, 108, 133)
CYAN = RGBColor(13, 138, 166)
BLUE = RGBColor(51, 102, 214)
GREEN = RGBColor(19, 155, 103)
MAGENTA = RGBColor(162, 75, 196)
ORANGE = RGBColor(215, 122, 11)
RED = RGBColor(197, 61, 74)
GRID = RGBColor(216, 226, 239)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=24,
    color=INK,
    bold=False,
    font="Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin=0.04,
    url=None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if url:
        run.hyperlink.address = url
    return box


def add_rich_text(slide, runs, x, y, w, h, *, size=20, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    for text, color, bold in runs:
        run = p.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_box(
    slide,
    x,
    y,
    w,
    h,
    *,
    fill=PANEL,
    line=GRID,
    radius=True,
    transparency=0,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_line(slide, x1, y1, x2, y2, *, color=GRID, width=2, arrow=False):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if arrow:
        line.line.end_arrowhead = True
    return line


def add_image_fit(slide, path, x, y, w, h, *, crop=False):
    path = str(path)
    with Image.open(path) as image:
        iw, ih = image.size
    image_ratio = iw / ih
    box_ratio = w / h
    if crop:
        pic = slide.shapes.add_picture(
            path, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        if image_ratio > box_ratio:
            shown = box_ratio / image_ratio
            pic.crop_left = pic.crop_right = (1 - shown) / 2
        else:
            shown = image_ratio / box_ratio
            pic.crop_top = pic.crop_bottom = (1 - shown) / 2
        return pic
    if image_ratio > box_ratio:
        fitted_w = w
        fitted_h = w / image_ratio
        x0, y0 = x, y + (h - fitted_h) / 2
    else:
        fitted_h = h
        fitted_w = h * image_ratio
        x0, y0 = x + (w - fitted_w) / 2, y
    return slide.shapes.add_picture(
        path, Inches(x0), Inches(y0), Inches(fitted_w), Inches(fitted_h)
    )


def add_title(slide, title, subtitle=None, section=None):
    if section:
        add_text(slide, section.upper(), 0.55, 0.23, 2.8, 0.28,
                 size=10, color=CYAN, bold=True)
    add_text(slide, title, 0.55, 0.52, 12.2, 0.62,
             size=27, color=INK, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.08, 12.0, 0.35,
                 size=12, color=MUTED)
    add_line(slide, 0.58, 1.43, 12.75, 1.43, color=GRID, width=1)


def add_footer(slide, number, source=None, source_url=None):
    add_text(slide, f"{number:02d}", 12.35, 7.12, 0.45, 0.22,
             size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    if source:
        add_text(slide, source, 0.55, 7.08, 10.8, 0.24,
                 size=8, color=MUTED, url=source_url)


def add_metric(slide, value, label, x, y, w, *, color=CYAN):
    add_box(slide, x, y, w, 1.15, fill=PANEL, line=GRID)
    add_text(slide, value, x + 0.12, y + 0.12, w - 0.24, 0.47,
             size=25, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.12, y + 0.63, w - 0.24, 0.32,
             size=10, color=MUTED, align=PP_ALIGN.CENTER)


def add_section_slide(number, kicker, title, subtitle, color):
    slide = prs.slides.add_slide(BLANK)
    add_bg(slide)
    add_text(slide, kicker.upper(), 0.75, 1.05, 4.5, 0.35,
             size=13, color=color, bold=True)
    add_text(slide, title, 0.75, 1.55, 11.6, 1.25,
             size=42, color=INK, bold=True)
    add_text(slide, subtitle, 0.78, 3.05, 9.8, 0.75,
             size=18, color=MUTED)
    add_line(slide, 0.78, 4.25, 11.75, 4.25, color=color, width=4)
    for i, width in enumerate((1.65, 2.3, 3.0, 3.8)):
        add_box(slide, 0.8 + i * 2.9, 5.0, width, 0.72,
                fill=PANEL_2, line=color)
        add_text(slide, str(i + 1).zfill(2), 0.95 + i * 2.9, 5.14,
                 0.5, 0.35, size=13, color=color, bold=True)
    add_footer(slide, number)
    return slide


# 01 — Title
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_box(slide, 0.0, 0.0, 4.8, 7.5, fill=PANEL_2, line=PANEL_2, radius=False)
add_text(slide, "vLLM · EPLB", 0.72, 0.65, 3.4, 0.35,
         size=13, color=CYAN, bold=True)
add_text(slide, "Contention-Aware\nExpert Migration\nBatching", 0.72, 1.35,
         11.8, 2.25, size=39, color=INK, bold=True,
         valign=MSO_ANCHOR.TOP)
add_text(slide, "From an optimization model to a merged production system",
         0.76, 4.2, 7.3, 0.55, size=19, color=MUTED)
add_box(slide, 8.35, 0.82, 4.15, 5.35, fill=WHITE, line=CYAN)
add_image_fit(slide, ILP / "example_hotspot_after.png", 8.48, 0.95, 3.89, 5.09)
add_text(slide, "PR #52641 · MERGED", 0.76, 6.35, 4.3, 0.35,
         size=14, color=GREEN, bold=True, url=PR_URL)
add_text(slide, "October 2026", 10.05, 6.45, 2.1, 0.3,
         size=12, color=MUTED, align=PP_ALIGN.RIGHT)
add_footer(slide, 1)


# 02 — Agenda
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "Roadmap", "Two questions: can contention be optimized, and does it help real serving?")
items = [
    ("01", "Upstream impact", "The merged vLLM contribution", GREEN),
    ("02", "ILP model", "Theoretical optimum and trade-off", CYAN),
    ("03", "E2E experiment", "Serving quality and NIC traffic", BLUE),
    ("04", "Profiling", "Scheduler cost versus time saved", MAGENTA),
]
for i, (num, title, sub, color) in enumerate(items):
    x = 0.75 + (i % 2) * 6.2
    y = 1.82 + (i // 2) * 2.05
    add_box(slide, x, y, 5.7, 1.55, fill=PANEL, line=color)
    add_text(slide, num, x + 0.25, y + 0.2, 0.7, 0.4,
             size=17, color=color, bold=True)
    add_text(slide, title, x + 1.08, y + 0.18, 4.25, 0.42,
             size=20, color=INK, bold=True)
    add_text(slide, sub, x + 1.08, y + 0.72, 4.25, 0.38,
             size=12, color=MUTED)
add_footer(slide, 2)


# 03 — Merged PR
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "The work is now part of vLLM", "A production contribution, not only a simulation")
add_box(slide, 0.58, 1.7, 9.15, 4.95, fill=PANEL, line=GREEN)
pic = add_image_fit(slide, HERE / "github_pr_52641.png", 0.72, 1.84, 8.87, 4.67, crop=True)
pic.click_action.hyperlink.address = PR_URL
add_metric(slide, "MERGED", "into vllm-project/vllm:main", 10.05, 1.85, 2.55, color=GREEN)
add_metric(slide, "44", "commits reviewed", 10.05, 3.25, 2.55, color=CYAN)
add_metric(slide, "12/12", "scheduler tests passed", 10.05, 4.65, 2.55, color=BLUE)
add_text(slide, "github.com/vllm-project/vllm/pull/52641", 9.92, 6.22, 2.82, 0.34,
         size=9, color=CYAN, align=PP_ALIGN.CENTER, url=PR_URL)
add_footer(slide, 3, "Source: vLLM PR #52641", PR_URL)


# 04 — Problem
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "The problem: migration bursts collide with inference",
          "Async EPLB moves expert weights while user requests are still running")
add_box(slide, 0.62, 1.72, 12.08, 3.18, fill=WHITE, line=RED)
add_image_fit(slide, ILP / "example_hotspot_before.png", 0.74, 1.84, 11.84, 2.94)
effects = [
    ("Shared ranks", "Multiple transfers compete at the same endpoint", RED),
    ("Shared NICs", "Cross-server flows create short traffic bursts", ORANGE),
    ("Serving impact", "Less network headroom for active requests", BLUE),
]
for i, (name, desc, color) in enumerate(effects):
    x = 0.72 + i * 4.12
    add_box(slide, x, 5.28, 3.72, 1.08, fill=PANEL, line=color)
    add_text(slide, name, x + 0.18, 5.43, 1.28, 0.3,
             size=13, color=color, bold=True)
    add_text(slide, desc, x + 1.42, 5.38, 2.08, 0.46,
             size=10, color=INK)
add_footer(slide, 4, "PR summary: contention-aware expert migration batching", PR_URL)


# 05 — ILP divider
add_section_slide(5, "Part I", "ILP model", "What is the best possible batching schedule under the production constraints?", CYAN)


# 06 — ILP model with requested two images
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "The model connects batch assignment to physical links",
          "Two core equations explain the scheduling rule and the optimization target", "Part I · ILP")
add_box(slide, 0.5, 1.68, 2.68, 5.15, fill=WHITE, line=CYAN)
add_image_fit(slide, ILP / "example_hotspot_after.png", 0.59, 1.77, 2.5, 4.97)
add_box(slide, 3.35, 1.68, 3.02, 5.15, fill=WHITE, line=MAGENTA)
add_image_fit(slide, ILP / "example_hotspot_links.png", 3.44, 1.77, 2.84, 4.97)

equations = [
    (
        "RANK-DISJOINT BATCH",
        "Σₚ:ᵣ∈ₚ xₚ,ᵦ ≤ 1",
        "In one batch, each rank communicates with at most one peer.",
        MAGENTA,
    ),
    (
        "WORST-LINK CONTENTION",
        "U = maxₑ,ᵦ (ρₑ + Lₑ,ᵦ / Cₑ)",
        "Peak load = inference traffic + scheduled migration traffic.",
        ORANGE,
    ),
]
for i, (tag, formula, meaning, color) in enumerate(equations):
    y = 1.82 + i * 1.62
    add_box(slide, 6.58, y, 6.08, 1.35, fill=PANEL, line=color)
    add_text(slide, tag, 6.84, y + 0.12, 2.15, 0.24,
             size=9, color=color, bold=True)
    add_text(slide, formula, 7.02, y + 0.43, 5.2, 0.38,
             size=20, color=INK, bold=True, font="Cambria Math",
             align=PP_ALIGN.CENTER)
    add_text(slide, meaning, 6.84, y + 0.94, 5.55, 0.25,
             size=10, color=MUTED, align=PP_ALIGN.CENTER)
add_box(slide, 6.58, 5.06, 6.08, 0.92, fill=PANEL_2, line=CYAN)
add_rich_text(slide, [
    ("ρₑ", MAGENTA, True),
    (" is existing inference load;  ", MUTED, False),
    ("Lₑ,ᵦ / Cₑ", ORANGE, True),
    (" is the migration share.\nMinimizing ", MUTED, False),
    ("U", CYAN, True),
    (" spreads transfers away from the busiest shared links.", MUTED, False),
], 6.84, 5.18, 5.56, 0.65, size=11, align=PP_ALIGN.CENTER)
add_box(slide, 6.58, 6.18, 6.08, 0.65, fill=PANEL_2, line=GREEN)
add_rich_text(slide, [
    ("Lexicographic:  ", MUTED, False),
    ("min batches", GREEN, True),
    ("  →  ", MUTED, False),
    ("min U", CYAN, True),
    ("  →  ", MUTED, False),
    ("min transfer time", INK, True),
], 6.78, 6.31, 5.68, 0.31, size=12, align=PP_ALIGN.CENTER)
add_footer(slide, 6, "Source: ILP/README.md, example_hotspot_after.png, example_hotspot_links.png")


# 07 — ILP result
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "The optimum reduces peak contention—not migration time",
          "Sequential batches deliberately trade completion time for serving headroom", "Part I · ILP")
add_text(slide, "PEAK TOTAL LINK UTILIZATION", 0.78, 1.75, 4.8, 0.3,
         size=11, color=CYAN, bold=True)
max_h = 3.05
for x, value, label, color in [
    (1.05, 2.10, "Unbatched", RED),
    (3.35, 1.10, "ILP optimum", GREEN),
]:
    height = max_h * value / 2.3
    add_box(slide, x, 5.35 - height, 1.35, height,
            fill=color, line=color, radius=False)
    add_text(slide, f"{value:.2f}", x - 0.08, 5.04 - height, 1.5, 0.35,
             size=19, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x - 0.2, 5.55, 1.75, 0.35,
             size=11, color=INK, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "−47.6%", 2.2, 3.15, 1.15, 0.4,
         size=18, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

add_text(slide, "MIGRATION COMPLETION TIME", 6.55, 1.75, 4.8, 0.3,
         size=11, color=ORANGE, bold=True)
for x, value, label, color in [
    (6.95, 1.55, "Lower bound", CYAN),
    (9.25, 1.755556, "4 batches", ORANGE),
]:
    height = max_h * value / 2.0
    add_box(slide, x, 5.35 - height, 1.35, height,
            fill=color, line=color, radius=False)
    add_text(slide, f"{value:.3f}", x - 0.08, 5.04 - height, 1.5, 0.35,
             size=19, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x - 0.2, 5.55, 1.75, 0.35,
             size=11, color=INK, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "+13.3%", 8.1, 3.0, 1.2, 0.4,
         size=18, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)

add_box(slide, 0.9, 6.23, 11.5, 0.55, fill=PANEL_2, line=CYAN)
add_text(slide,
         "M_batch = Σᵦ Tᵦ + δΣᵦ yᵦ  ≥  M_unbatched   ·   Lower U means less interference, not faster migration",
         1.13, 6.34, 11.02, 0.3, size=13, color=INK,
         font="Cambria Math", align=PP_ALIGN.CENTER)
add_footer(slide, 7, "Computed from ILP/example_hotspot.json")


# 08 — Experiment divider
add_section_slide(8, "Part II", "Real-system experiments", "E2E serving validates the benefit; Nsight profiling measures the cost.", GREEN)


# 09 — E2E setup
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "E2E setup: four-node async serving over 10 GbE",
          "Fresh server for each case · identical workload · batching is the only changed switch", "Part II · E2E")
add_box(slide, 0.62, 1.72, 8.0, 4.7, fill=PANEL, line=BLUE)
for i in range(4):
    x = 1.03 + i * 1.82
    add_box(slide, x, 2.18, 1.35, 1.33, fill=PANEL_2, line=CYAN)
    add_text(slide, f"NODE {i}", x + 0.08, 2.3, 1.18, 0.28,
             size=11, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "RTX 6000\n24 GB", x + 0.08, 2.7, 1.18, 0.55,
             size=13, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, x + 0.68, 3.52, x + 0.68, 4.3,
             color=GREEN, width=3)
add_line(slide, 1.68, 4.3, 7.16, 4.3, color=GREEN, width=5)
add_text(slide, "10 GbE · NIXL 1.3.2 · no RDMA", 2.65, 4.48, 3.55, 0.35,
         size=13, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "Qwen3-30B-A3B-Instruct-2507", 1.15, 5.32, 3.0, 0.35,
         size=13, color=INK, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "TP = EP = 4", 4.25, 5.32, 1.4, 0.35,
         size=13, color=INK, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "async EPLB", 6.05, 5.32, 1.4, 0.35,
         size=13, color=INK, bold=True, align=PP_ALIGN.CENTER)

setup = [
    ("500", "requests"),
    ("32", "max concurrency"),
    ("300", "output tokens"),
    ("50", "EPLB step interval"),
]
for i, (value, label) in enumerate(setup):
    y = 1.82 + i * 1.18
    add_metric(slide, value, label, 9.05, y, 3.2,
               color=(CYAN, BLUE, MAGENTA, GREEN)[i])
add_box(slide, 9.05, 6.42, 3.2, 0.46, fill=PANEL_2, line=GREEN)
add_text(slide, "500 / 500 completed in both runs", 9.16, 6.51, 2.98, 0.24,
         size=10, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
add_footer(slide, 9, "Source: serving_nixl_20260911_async_random500/COMMANDS.md")


# 10 — E2E result
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "Batching improves every reported serving metric",
          "Positive values mean higher throughput or lower latency", "Part II · E2E")
add_box(slide, 0.56, 1.72, 12.2, 3.95, fill=WHITE, line=GREEN)
add_image_fit(slide, E2E / "e2e_improvement.png", 0.7, 1.86, 11.92, 3.67)
add_metric(slide, "+2.28%", "output throughput", 1.0, 5.82, 2.65, color=BLUE)
add_metric(slide, "−3.18%", "E2EL P50", 3.96, 5.82, 2.65, color=GREEN)
add_metric(slide, "−1.30%", "TTFT P99", 6.92, 5.82, 2.65, color=CYAN)
add_metric(slide, "−0.87%", "TPOT P99", 9.88, 5.82, 2.65, color=MAGENTA)
add_footer(slide, 10, "Source: e2e_summary.csv and e2e_improvement.png")


# 11 — NIC traffic
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "The mechanism is visible at the NIC",
          "One-second Linux RX+TX samples on the head node", "Part II · E2E")
add_box(slide, 0.56, 1.65, 10.15, 5.28, fill=WHITE, line=CYAN)
add_image_fit(slide, E2E / "nic_timeseries.png", 0.68, 1.76, 9.92, 5.04)
add_metric(slide, "599.81", "MB/s P99 · batching off", 10.95, 1.92, 1.93, color=RED)
add_metric(slide, "538.40", "MB/s P99 · batching on", 10.95, 3.42, 1.93, color=GREEN)
add_metric(slide, "−10.24%", "head-node NIC P99", 10.95, 4.92, 1.93, color=CYAN)
add_text(slide, "Less bursty migration traffic leaves more headroom for inference.",
         10.92, 6.32, 1.98, 0.54, size=10, color=INK,
         bold=True, align=PP_ALIGN.CENTER)
add_footer(slide, 11, "Source: nic.tsv, nic_summary.csv and nic_timeseries.png")


# 12 — Profiling method
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "Profiling isolates the scheduler—not the data transfer",
          "NVTX-only Nsight Systems capture on all four ranks", "Part II · Profiling")
add_box(slide, 0.65, 1.8, 12.0, 2.15, fill=PANEL, line=MAGENTA)
labels = ["Rebalance cycle", "Schedule all 48 layers", "Layer-by-layer transfer", "Next cycle"]
colors = [BLUE, MAGENTA, GREEN, BLUE]
widths = [1.7, 2.7, 4.15, 1.55]
x = 1.0
for i, (label, color, width) in enumerate(zip(labels, colors, widths)):
    add_box(slide, x, 2.45, width, 0.58, fill=color, line=color, radius=False)
    add_text(slide, label, x + 0.08, 2.54, width - 0.16, 0.33,
             size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    if i < len(labels) - 1:
        add_line(slide, x + width, 2.74, x + width + 0.3, 2.74,
                 color=MUTED, width=2, arrow=True)
    x += width + 0.35
add_text(slide, "Nsight range: “eplb: schedule migration batches”",
         3.8, 3.28, 5.8, 0.35, size=13, color=MAGENTA,
         bold=True, align=PP_ALIGN.CENTER)

cards = [
    ("4", "rank traces", CYAN),
    ("34", "scheduler calls / rank", MAGENTA),
    ("48", "MoE layers / call", GREEN),
    ("7.0–7.8 ms", "median cost / call", ORANGE),
]
for i, (value, label, color) in enumerate(cards):
    add_metric(slide, value, label, 0.75 + i * 3.1, 4.48, 2.65, color=color)
add_box(slide, 1.65, 6.05, 10.0, 0.62, fill=PANEL_2, line=CYAN)
add_text(slide,
         "Same model and workload as E2E; only NVTX capture was added. The four ranks schedule concurrently.",
         1.9, 6.17, 9.5, 0.34, size=12, color=INK,
         align=PP_ALIGN.CENTER)
add_footer(slide, 12, "Source: profiling COMMANDS.md and trace/rank0.nsys-rep … rank3.nsys-rep")


# 13 — Per-rank profile
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "Scheduler overhead is small and balanced across ranks",
          "34 concurrent cycle-wide calls per rank", "Part II · Profiling")
values = [302.081, 253.152, 255.239, 241.190]
p50s = [7.819, 7.090, 7.035, 7.005]
base_y = 5.62
for i, (total, p50) in enumerate(zip(values, p50s)):
    x = 1.25 + i * 2.85
    height = 3.55 * total / 330
    color = (MAGENTA, CYAN, BLUE, GREEN)[i]
    add_box(slide, x, base_y - height, 1.35, height,
            fill=color, line=color, radius=False)
    add_text(slide, f"{total:.1f} ms", x - 0.18, base_y - height - 0.42,
             1.7, 0.35, size=16, color=color, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(slide, f"Rank {i}", x - 0.05, base_y + 0.12, 1.45, 0.35,
             size=12, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"P50 {p50:.3f} ms", x - 0.2, base_y + 0.53, 1.75, 0.28,
             size=10, color=MUTED, align=PP_ALIGN.CENTER)
add_line(slide, 0.88, base_y, 12.45, base_y, color=GRID, width=2)
add_box(slide, 9.95, 1.62, 2.45, 0.75, fill=PANEL_2, line=GREEN)
add_text(slide, "Worst total: 302.081 ms", 10.1, 1.79, 2.15, 0.34,
         size=12, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
add_footer(slide, 13, "Source: profile_summary.csv")


# 14 — Cost versus saved
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_title(slide, "The optimization pays for itself by a wide margin",
          "Profiling cost is compared with the independent clean E2E A/B run", "Part II · Profiling")
add_box(slide, 0.62, 1.68, 9.35, 5.18, fill=WHITE, line=GREEN)
add_image_fit(slide, PROFILE / "scheduler_cost_vs_saved.png", 0.78, 1.83, 9.03, 4.88)
add_metric(slide, "302.081 ms", "worst-rank total scheduler cost", 10.3, 1.98, 2.43, color=ORANGE)
add_metric(slide, "105.932 s", "serving time saved", 10.3, 3.52, 2.43, color=GREEN)
add_metric(slide, "0.285%", "cost / saved time", 10.3, 5.06, 2.43, color=CYAN)
add_footer(slide, 14, "Sources: profile_summary.csv + E2E off/on duration")


# 15 — Takeaways
slide = prs.slides.add_slide(BLANK)
add_bg(slide)
add_text(slide, "TAKEAWAYS", 0.72, 0.72, 3.0, 0.35,
         size=13, color=GREEN, bold=True)
add_text(slide, "A model, a production implementation,\nand measured serving gains.",
         0.72, 1.28, 11.6, 1.25, size=34, color=INK, bold=True,
         valign=MSO_ANCHOR.TOP)
takeaways = [
    ("01", "Optimization", "ILP proves peak contention is schedulable under rank-disjoint constraints.", CYAN),
    ("02", "Production", "The opt-in scheduler is merged into upstream vLLM as PR #52641.", GREEN),
    ("03", "Evidence", "NIC P99 −10.24%, throughput +2.28%, E2EL P50 −3.18%.", BLUE),
    ("04", "Efficiency", "Worst-rank scheduler cost is only 0.285% of serving time saved.", MAGENTA),
]
for i, (num, title, desc, color) in enumerate(takeaways):
    x = 0.72 + (i % 2) * 6.18
    y = 3.05 + (i // 2) * 1.55
    add_box(slide, x, y, 5.62, 1.14, fill=PANEL, line=color)
    add_text(slide, num, x + 0.22, y + 0.18, 0.55, 0.3,
             size=13, color=color, bold=True)
    add_text(slide, title, x + 0.9, y + 0.14, 1.55, 0.34,
             size=15, color=INK, bold=True)
    add_text(slide, desc, x + 0.9, y + 0.54, 4.28, 0.4,
             size=10, color=MUTED)
add_text(slide, "github.com/vllm-project/vllm/pull/52641", 4.4, 6.55, 4.55, 0.35,
         size=13, color=CYAN, bold=True, align=PP_ALIGN.CENTER, url=PR_URL)
add_footer(slide, 15)


prs.core_properties.title = "Contention-Aware Expert Migration Batching"
prs.core_properties.subject = "ILP modeling and real-system experiments for vLLM PR #52641"
prs.core_properties.author = "DOCCA0"
prs.core_properties.keywords = "vLLM, EPLB, expert migration, ILP, NIXL, profiling"
prs.save(OUT)
print(OUT)
