"""Apply the final defense-oriented revisions to the existing presentation.

This script intentionally edits the current deck instead of rebuilding it, so
manual changes made in PowerPoint or WPS are retained outside the revised
slides.
"""

from __future__ import annotations

import tempfile
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
DECK = HERE / "eplb_migration_batching.pptx"
QWEN = HERE / "qwen3_moe_architecture.png"
ILP = ROOT / "benchmarks/eplb_rebalance/ILP"

NAVY = RGBColor(19, 48, 82)
BLUE = RGBColor(45, 103, 164)
LIGHT_BLUE = RGBColor(225, 237, 248)
CYAN = RGBColor(21, 142, 166)
GREEN = RGBColor(27, 141, 96)
LIGHT_GREEN = RGBColor(229, 245, 238)
ORANGE = RGBColor(221, 133, 42)
LIGHT_ORANGE = RGBColor(251, 239, 222)
RED = RGBColor(191, 65, 73)
LIGHT_RED = RGBColor(250, 230, 232)
PURPLE = RGBColor(129, 83, 160)
INK = RGBColor(28, 39, 53)
MUTED = RGBColor(93, 108, 126)
GRID = RGBColor(210, 220, 231)
BG = RGBColor(249, 251, 253)
WHITE = RGBColor(255, 255, 255)


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def clear_body(slide, keep_subtitle: bool = False) -> None:
    """Retain the background, section label, title, divider, and optional subtitle."""
    keep = {0, 1, 2, 3}
    if keep_subtitle:
        keep = {0, 1, 2, 3, 4}
    for index, shape in reversed(list(enumerate(slide.shapes))):
        if index not in keep:
            remove_shape(shape)


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill=WHITE,
    line=GRID,
    radius: bool = True,
    width: float = 1.2,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(width)
    return shape


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 20,
    color=INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.04)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color=GRID,
    width: float = 2,
):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_image_fit(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as image:
        ratio = image.width / image.height
    box_ratio = w / h
    if ratio > box_ratio:
        height = w / ratio
        return slide.shapes.add_picture(
            str(path), Inches(x), Inches(y + (h - height) / 2), width=Inches(w)
        )
    width = h * ratio
    return slide.shapes.add_picture(
        str(path), Inches(x + (w - width) / 2), Inches(y), height=Inches(h)
    )


def set_title(slide, text: str) -> None:
    slide.shapes[2].text_frame.paragraphs[0].runs[0].text = text


def remove_footers(prs: Presentation) -> None:
    for slide in prs.slides:
        for shape in reversed(list(slide.shapes)):
            if shape.top >= Inches(6.9):
                remove_shape(shape)


def revise_outline(slide) -> None:
    clear_body(slide)
    set_title(slide, "Outline")
    sections = [
        ("01", "Background", "MoE, EP, and asynchronous EPLB", BLUE),
        ("02", "Model", "Contention objective and offline ILP", PURPLE),
        ("03", "Design", "Rank-disjoint migration batching", ORANGE),
        ("04", "Evaluation", "E2E serving, NIC traffic, and profiling", GREEN),
    ]
    for index, (number, title, detail, color) in enumerate(sections):
        x = 0.85 + (index % 2) * 6.0
        y = 1.85 + (index // 2) * 2.15
        add_box(slide, x, y, 5.55, 1.55, fill=WHITE, line=color, width=1.7)
        add_text(
            slide,
            number,
            x + 0.24,
            y + 0.28,
            0.78,
            0.55,
            size=27,
            color=color,
            bold=True,
        )
        add_text(
            slide, title, x + 1.1, y + 0.2, 3.9, 0.46, size=24, color=NAVY, bold=True
        )
        add_text(slide, detail, x + 1.1, y + 0.72, 4.0, 0.5, size=15, color=MUTED)


def add_qwen_slide(prs: Presentation) -> None:
    if any(
        any(
            getattr(shape, "text", "") == "Qwen3 MoE Architecture"
            for shape in slide.shapes
        )
        for slide in prs.slides
    ):
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    add_text(slide, "BACKGROUND", 0.55, 0.18, 4.6, 0.25, size=11, color=BLUE, bold=True)
    add_text(
        slide,
        "Qwen3 MoE Architecture",
        0.55,
        0.48,
        12.1,
        0.55,
        size=27,
        color=NAVY,
        bold=True,
    )
    add_line(slide, 0.58, 1.4, 12.76, 1.4, color=GRID, width=1.5)
    add_image_fit(slide, QWEN, 0.82, 1.58, 11.7, 5.55)
    add_box(slide, 9.05, 5.82, 3.08, 0.68, fill=LIGHT_GREEN, line=GREEN)
    add_text(
        slide,
        "EPLB moves the expert weights inside each MoE layer",
        9.22,
        5.91,
        2.75,
        0.47,
        size=14,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    slide_id = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(slide_id)
    prs.slides._sldIdLst.insert(3, slide_id)


def revise_async(slide) -> None:
    clear_body(slide)
    set_title(slide, "Asynchronous EPLB")
    steps = [
        ("1", "Collect\nexpert load", BLUE),
        ("2", "Compute new\nplacement", CYAN),
        ("3", "Migrate\nexpert weights", ORANGE),
        ("4", "Activate at\nsafe boundary", GREEN),
    ]
    add_box(slide, 0.62, 1.68, 12.08, 2.0, fill=WHITE, line=BLUE, width=1.5)
    for index, (number, label, color) in enumerate(steps):
        x = 0.95 + index * 3.0
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(2.08), Inches(0.55), Inches(0.55)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        add_text(
            slide,
            number,
            x,
            2.09,
            0.55,
            0.53,
            size=15,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide, label, x + 0.75, 1.96, 1.8, 0.74, size=18, color=NAVY, bold=True
        )
        if index < 3:
            add_line(slide, x + 2.45, 2.36, x + 2.8, 2.36, color=MUTED, width=2.3)
    add_text(
        slide,
        "Foreground inference continues during migration",
        3.2,
        3.1,
        6.9,
        0.34,
        size=16,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 0.62, 4.05, 5.75, 2.12, fill=LIGHT_RED, line=RED, width=1.5)
    add_text(
        slide,
        "UNBATCHED MIGRATION",
        0.92,
        4.25,
        2.55,
        0.3,
        size=14,
        color=RED,
        bold=True,
    )
    # A converging endpoint diagram replaces the ambiguous parallel lines.
    for i, color in enumerate((RED, ORANGE, PURPLE)):
        y = 4.85 + i * 0.42
        add_text(
            slide,
            f"R{i}",
            1.0,
            y - 0.12,
            0.55,
            0.32,
            size=13,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_line(slide, 1.58, y, 4.55, 5.28, color=color, width=3.0)
    endpoint = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(4.55), Inches(4.88), Inches(0.78), Inches(0.78)
    )
    endpoint.fill.solid()
    endpoint.fill.fore_color.rgb = RED
    endpoint.line.fill.background()
    add_text(
        slide,
        "NIC",
        4.55,
        4.9,
        0.78,
        0.72,
        size=13,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Concurrent flows converge on one endpoint → burst",
        1.05,
        5.86,
        4.85,
        0.22,
        size=11,
        color=RED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 6.67, 4.05, 6.03, 2.12, fill=LIGHT_GREEN, line=GREEN, width=1.5)
    add_text(slide, "GOAL", 6.98, 4.25, 1.0, 0.3, size=14, color=GREEN, bold=True)
    add_text(
        slide,
        "Reshape migration traffic",
        7.22,
        4.73,
        4.95,
        0.48,
        size=24,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Protect inference without changing expert placement",
        7.18,
        5.35,
        5.05,
        0.36,
        size=15,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def revise_problem(slide) -> None:
    clear_body(slide)
    set_title(slide, "Problem Statement")
    # Layout mirrors the reference presentation: one root, four branches,
    # icon blocks, and the corresponding problem definition on the right.
    add_text(
        slide,
        "Migration\nscheduling",
        0.0,
        3.15,
        1.9,
        0.72,
        size=18,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    root = add_box(
        slide, 0.42, 4.0, 1.12, 1.12, fill=NAVY, line=GREEN, radius=False, width=1.8
    )
    del root
    add_text(
        slide,
        "↔",
        0.42,
        4.0,
        1.12,
        1.05,
        size=42,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_line(slide, 1.54, 4.55, 2.38, 4.55, color=BLUE, width=1.5)
    add_line(slide, 2.38, 2.02, 2.38, 6.02, color=BLUE, width=1.5)
    rows = [
        ("Given", "G", "Old/new placements; topology and residual link capacity"),
        ("Decide", "D", "Assign each directed migration flow to one batch"),
        ("Objective", "O", "Minimize worst link pressure seen by inference"),
        (
            "Constraints",
            "C",
            "Each flow once; at most one incident flow per rank and batch",
        ),
    ]
    for i, (label, symbol, description) in enumerate(rows):
        y = 1.72 + i * 1.32
        add_text(
            slide,
            label,
            2.45,
            y + 0.08,
            1.42,
            0.45,
            size=18,
            color=INK,
            align=PP_ALIGN.CENTER,
        )
        add_line(slide, 2.38, y + 0.55, 4.18, y + 0.55, color=BLUE, width=1.5)
        icon = add_box(
            slide, 4.18, y, 0.92, 0.92, fill=NAVY, line=GREEN, radius=False, width=1.8
        )
        del icon
        add_text(
            slide,
            symbol,
            4.18,
            y,
            0.92,
            0.86,
            size=27,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide, "▪  " + description, 5.4, y - 0.02, 7.25, 0.95, size=18, color=INK
        )


def revise_ilp_example(slide) -> None:
    clear_body(slide)
    set_title(slide, "ILP Hotspot Example")
    add_text(
        slide,
        "Eight ranks, four servers, and a shared inter-server bottleneck",
        0.57,
        1.02,
        12.0,
        0.32,
        size=12,
        color=MUTED,
    )
    panels = [
        (
            0.4,
            "TWO-TIER LINKS",
            ILP / "example_hotspot_links.png",
            BLUE,
            "Shared inter-server bottleneck",
        ),
        (
            4.48,
            "UNBATCHED",
            ILP / "example_hotspot_before.png",
            RED,
            "Peak 1.50  ·  Inference 4.75",
        ),
        (
            8.56,
            "ILP · 4 BATCHES",
            ILP / "example_hotspot_after.png",
            GREEN,
            "Peak 0.50  ·  Inference 2.25",
        ),
    ]
    for x, heading, image, color, caption in panels:
        add_box(slide, x, 1.67, 3.86, 4.9, fill=WHITE, line=color, width=1.5)
        add_text(
            slide,
            heading,
            x + 0.22,
            1.83,
            3.42,
            0.28,
            size=13,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_image_fit(slide, image, x + 0.32, 2.18, 3.22, 3.65)
        add_text(
            slide,
            caption,
            x + 0.22,
            6.0,
            3.42,
            0.31,
            size=13,
            color=color,
            bold=True,
            align=PP_ALIGN.CENTER,
        )


def revise_online(slide) -> None:
    clear_body(slide, keep_subtitle=True)
    set_title(slide, "Online Design")
    slide.shapes[3].text_frame.paragraphs[0].runs[
        0
    ].text = "The offline model yields one deployable invariant: rank-disjoint batches"
    add_box(slide, 0.7, 1.82, 5.2, 3.85, fill=LIGHT_BLUE, line=BLUE, width=1.5)
    add_text(
        slide, "OFFLINE ILP", 1.02, 2.08, 1.65, 0.28, size=14, color=BLUE, bold=True
    )
    add_text(
        slide, "Exact reference", 1.02, 2.55, 3.8, 0.45, size=26, color=NAVY, bold=True
    )
    add_text(
        slide,
        "Explains the contention objective\nand validates small instances",
        1.02,
        3.2,
        4.15,
        0.78,
        size=18,
        color=INK,
    )
    add_text(
        slide,
        "Not used in serving",
        1.02,
        4.58,
        2.55,
        0.38,
        size=16,
        color=RED,
        bold=True,
    )
    add_line(slide, 6.15, 3.72, 7.05, 3.72, color=ORANGE, width=4)
    add_text(
        slide,
        "invariant",
        6.18,
        3.28,
        0.85,
        0.3,
        size=12,
        color=ORANGE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 7.3, 1.82, 5.3, 3.85, fill=LIGHT_GREEN, line=GREEN, width=1.5)
    add_text(
        slide,
        "ONLINE SCHEDULER",
        7.62,
        2.08,
        2.1,
        0.28,
        size=14,
        color=GREEN,
        bold=True,
    )
    add_text(
        slide, "Greedy first-fit", 7.62, 2.55, 3.8, 0.45, size=26, color=NAVY, bold=True
    )
    add_text(
        slide,
        "Coalesce by rank pair\nGroup endpoint-disjoint flows",
        7.62,
        3.2,
        4.15,
        0.78,
        size=18,
        color=INK,
    )
    add_text(
        slide,
        "Deterministic · lightweight · opt-in",
        7.62,
        4.58,
        4.1,
        0.38,
        size=16,
        color=GREEN,
        bold=True,
    )
    add_box(slide, 2.0, 6.03, 9.35, 0.52, fill=WHITE, line=GRID)
    add_text(
        slide,
        "Placement semantics remain unchanged; only migration order changes.",
        2.15,
        6.08,
        9.05,
        0.36,
        size=15,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def rename_titles(prs: Presentation) -> None:
    replacements = {
        "The correct objective depends on the execution mode": "Serving Objective",
        "ILP formulation exposes the contention trade-off": "ILP Formulation",
        "Online scheduler: five transformations from placement to batches": (
            "Scheduling Pipeline"
        ),
        "Greedy first-fit groups independent flows in the same batch": (
            "Greedy First-Fit"
        ),
        "Correctness and cost are explicit—not empirical assumptions": (
            "Correctness and Complexity"
        ),
        "vLLM separates cycle-wide scheduling from layer-wise transfer": (
            "vLLM Integration"
        ),
        "Controlled A/B evaluation on four-node cluster": "Experimental Setup",
        "Controlled A/B evaluation on a four-node cluster": "Experimental Setup",
        "End-to-end results: every reported serving metric improves": "E2E Results",
        "The benefit is consistent across throughput and latency": (
            "Serving Improvements"
        ),
        "Lower NIC tail traffic supports proposed mechanism": "NIC Traffic",
        "Lower NIC tail traffic supports the proposed mechanism": "NIC Traffic",
        "Profiling measures scheduler cost in serving context": "Profiling Method",
        "Profiling measures scheduler cost in the serving context": "Profiling Method",
        "Scheduler cost is 0.285% of serving time saved": "Profiling Results",
    }
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if shape.text in replacements:
                shape.text_frame.paragraphs[0].runs[0].text = replacements[shape.text]


def find_slide(prs: Presentation, *titles: str):
    for slide in prs.slides:
        texts = {
            shape.text for shape in slide.shapes if shape.has_text_frame and shape.text
        }
        if texts.intersection(titles):
            return slide
    raise ValueError(f"Could not find slide with any title in {titles}")


def clean_office_math(path: Path) -> None:
    """Remove invisible Pandoc spacing glyphs while preserving editable OMML."""
    substitutions = {
        "\u200b": "",
        "\u2001": " ",
        "\u2004": " ",
        "\u2005": " ",
        "\u2009": " ",
    }
    with tempfile.TemporaryDirectory() as temp_dir:
        temp = Path(temp_dir) / path.name
        with (
            zipfile.ZipFile(path, "r") as source,
            zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target,
        ):
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename.startswith("ppt/slides/") and item.filename.endswith(
                    ".xml"
                ):
                    text = data.decode("utf-8")
                    for old, new in substitutions.items():
                        text = text.replace(old, new)
                    data = text.encode("utf-8")
                target.writestr(item, data)
        path.write_bytes(temp.read_bytes())


def main() -> None:
    prs = Presentation(DECK)
    outline = find_slide(prs, "Roadmap and research questions", "Outline")
    async_eplb = find_slide(
        prs,
        "Async EPLB creates network interference during migration",
        "Asynchronous EPLB",
    )
    problem = find_slide(
        prs,
        "Problem statement: schedule flows, not individual tensors",
        "Problem Statement",
    )
    ilp_example = find_slide(
        prs,
        "The constrained optimum reduces contention, not migration time",
        "ILP Hotspot Example",
    )
    online = find_slide(
        prs,
        "From an exact model to a production scheduler",
        "Online Design",
    )
    remove_footers(prs)
    rename_titles(prs)
    revise_outline(outline)
    revise_async(async_eplb)
    revise_problem(problem)
    revise_ilp_example(ilp_example)
    revise_online(online)
    add_qwen_slide(prs)
    prs.save(DECK)
    clean_office_math(DECK)


if __name__ == "__main__":
    main()
