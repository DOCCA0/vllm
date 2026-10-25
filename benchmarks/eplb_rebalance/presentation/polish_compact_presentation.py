"""Polish the manually shortened defense deck without rebuilding it."""

from __future__ import annotations

import copy
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

DECK = Path(__file__).with_name("eplb_migration_batching.pptx")
ILP_LINK_IMAGE = DECK.parent.parent / "ILP" / "example_hotspot_links.png"


def OxmlElement(tag: str) -> etree._Element:  # noqa: N802
    """Create a namespaced DrawingML or PresentationML element."""
    return etree.Element(qn(tag))


def load_equations(formulas: dict[str, str]) -> dict[str, etree._Element]:
    """Convert LaTeX to editable Office Math through Pandoc."""
    markdown: list[str] = []
    for key, formula in formulas.items():
        markdown.extend([f"# {key}", "", f"$${formula}$$", "", "---", ""])

    with tempfile.TemporaryDirectory() as temp_dir:
        source = Path(temp_dir) / "equations.md"
        deck = Path(temp_dir) / "equations.pptx"
        source.write_text("\n".join(markdown), encoding="utf-8")
        subprocess.run(["pandoc", str(source), "-o", str(deck)], check=True)
        equations: dict[str, etree._Element] = {}
        namespace = {"a14": "http://schemas.microsoft.com/office/drawing/2010/main"}
        with zipfile.ZipFile(deck) as archive:
            for index, key in enumerate(formulas, start=1):
                root = etree.fromstring(archive.read(f"ppt/slides/slide{index}.xml"))
                math = root.find(".//a14:m", namespaces=namespace)
                if math is None:
                    raise RuntimeError(f"Pandoc did not generate Office Math: {key}")
                equations[key] = copy.deepcopy(math)
    return equations


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def ungroup_shape(group) -> None:
    """Restore a generated group to its original sibling shapes."""
    group_element = group._element
    parent = group_element.getparent()
    insertion_index = parent.index(group_element)
    for child in list(group_element):
        if child.tag in {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}:
            continue
        group_element.remove(child)
        parent.insert(insertion_index, child)
        insertion_index += 1
    parent.remove(group_element)


def remove_page_subtitles(prs: Presentation) -> int:
    """Remove the single-line explanatory text directly below page titles."""
    removed = 0
    for slide in prs.slides:
        for shape in reversed(list(slide.shapes)):
            if not shape.has_text_frame or not shape.text.strip():
                continue
            x = shape.left / Inches(1)
            y = shape.top / Inches(1)
            width = shape.width / Inches(1)
            if 0.45 <= x <= 0.7 and 0.95 <= y <= 1.12 and width >= 8:
                remove_shape(shape)
                removed += 1
    return removed


def set_shape_font_size(shape, size: float) -> None:
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.size = Pt(size)
        for run in paragraph.runs:
            run.font.size = Pt(size)


def set_text(
    shape,
    text: str,
    size: float,
    color: RGBColor,
    *,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.14)
    frame.margin_right = Inches(0.14)
    frame.margin_top = Inches(0.08)
    frame.margin_bottom = Inches(0.06)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_equation(
    slide,
    math: etree._Element,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: int,
    color: RGBColor,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    tx_body = box._element.txBody
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    for paragraph in list(tx_body.findall(f"{{{drawing_ns}}}p")):
        tx_body.remove(paragraph)

    paragraph = etree.Element(f"{{{drawing_ns}}}p")
    paragraph_properties = etree.SubElement(paragraph, f"{{{drawing_ns}}}pPr")
    paragraph_properties.set("algn", "ctr")
    equation = copy.deepcopy(math)
    for run in equation.findall(f".//{{{math_ns}}}r"):
        properties = etree.Element(f"{{{drawing_ns}}}rPr")
        properties.set("lang", "en-US")
        properties.set("sz", str(size * 100))
        properties.set("dirty", "0")
        fill = etree.SubElement(properties, f"{{{drawing_ns}}}solidFill")
        rgb = etree.SubElement(fill, f"{{{drawing_ns}}}srgbClr")
        rgb.set("val", "".join(f"{channel:02X}" for channel in color))
        latin = etree.SubElement(properties, f"{{{drawing_ns}}}latin")
        latin.set("typeface", "Cambria Math")
        text = run.find(f"{{{math_ns}}}t")
        run.insert(run.index(text) if text is not None else len(run), properties)
    paragraph.append(equation)
    tx_body.append(paragraph)
    body_properties = tx_body.find(f"{{{drawing_ns}}}bodyPr")
    body_properties.set("anchor", "ctr")
    body_properties.set("wrap", "none")


def add_panel(slide, x: float, y: float, width: float, height: float, color) -> None:
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    panel.line.color.rgb = color
    panel.line.width = Pt(1.4)


def rebuild_ilp_formulation(prs: Presentation) -> None:
    """Present the README model with consistent symbols and editable math."""
    formulas = load_equations(
        {
            "objective": r"\min T_{\mathrm{infer}}",
            "assign": r"\sum_{b\in B}x_{p,b}=1,\qquad \forall p\in P",
            "endpoint": (r"\sum_{p:r\in p}x_{p,b}\le 1,\qquad \forall r,\;b"),
            "activation": (r"x_{p,b}\le y_b,\qquad y_b\ge y_{b+1}"),
            "load": r"L_{e,b}=\sum_{p\in P}w_p f_{p,e}x_{p,b}",
            "slowdown": (
                r"T_{\mathrm{infer}}\ge T_0\left(1+"
                r"\frac{L_{e,b}}{C_e(1-\rho)}\right),"
                r"\quad \forall e,b"
            ),
            "x_variable": r"x_{p,b}\in\{0,1\}",
            "y_variable": r"y_b\in\{0,1\}",
            "time_variable": r"T_{\mathrm{infer}}\ge 0",
            "flow_size": r"w_p",
            "route_fraction": r"f_{p,e}",
            "capacity": r"C_e",
            "baseline_inputs": r"T_0,\ \rho",
        }
    )
    for slide in prs.slides:
        if not any(
            shape.has_text_frame and shape.text.strip() == "ILP Formulation"
            for shape in slide.shapes
        ):
            continue
        for shape in reversed(list(slide.shapes)[3:]):
            remove_shape(shape)

        navy = RGBColor(20, 52, 91)
        ink = RGBColor(31, 48, 70)
        blue = RGBColor(47, 111, 174)
        purple = RGBColor(125, 72, 161)
        orange = RGBColor(230, 126, 34)
        green = RGBColor(19, 143, 95)

        add_panel(slide, 0.62, 1.62, 3.18, 4.96, blue)
        heading = slide.shapes.add_textbox(
            Inches(0.86), Inches(1.78), Inches(2.7), Inches(0.28)
        )
        set_text(heading, "VARIABLES & INPUTS", 12, blue, bold=True)

        decision_heading = slide.shapes.add_textbox(
            Inches(0.86), Inches(2.14), Inches(2.7), Inches(0.25)
        )
        set_text(decision_heading, "DECISION VARIABLES", 10, purple, bold=True)
        decisions = [
            ("x_variable", "flow assignment"),
            ("y_variable", "active-batch flag"),
            ("time_variable", "worst inference time"),
        ]
        for index, (equation_key, meaning) in enumerate(decisions):
            y = 2.46 + index * 0.57
            add_equation(
                slide,
                formulas[equation_key],
                0.73,
                y,
                1.48,
                0.36,
                size=13,
                color=navy,
            )
            meaning_box = slide.shapes.add_textbox(
                Inches(2.08), Inches(y - 0.01), Inches(1.47), Inches(0.4)
            )
            set_text(meaning_box, meaning, 10.8, ink)

        input_heading = slide.shapes.add_textbox(
            Inches(0.86), Inches(4.2), Inches(2.7), Inches(0.25)
        )
        set_text(input_heading, "KNOWN INPUTS", 10, orange, bold=True)
        inputs = [
            ("flow_size", "experts in flow p"),
            ("route_fraction", "fraction routed on link e"),
            ("capacity", "capacity of link e"),
            ("baseline_inputs", "baseline time; inference share"),
        ]
        for index, (equation_key, meaning) in enumerate(inputs):
            y = 4.5 + index * 0.38
            add_equation(
                slide,
                formulas[equation_key],
                0.82,
                y,
                1.05,
                0.3,
                size=13,
                color=navy,
            )
            meaning_box = slide.shapes.add_textbox(
                Inches(1.75), Inches(y), Inches(1.8), Inches(0.32)
            )
            set_text(meaning_box, meaning, 9.8, ink)

        tie_break = slide.shapes.add_textbox(
            Inches(0.82), Inches(6.05), Inches(2.76), Inches(0.32)
        )
        set_text(
            tie_break,
            "Tie-break: minimize active batches after fixing Tᵢₙfₑᵣ",
            9.5,
            orange,
            bold=True,
            align=PP_ALIGN.CENTER,
        )

        add_panel(slide, 4.08, 1.62, 8.5, 0.72, green)
        objective_label = slide.shapes.add_textbox(
            Inches(4.3), Inches(1.78), Inches(2.15), Inches(0.3)
        )
        set_text(objective_label, "PRIMARY OBJECTIVE", 11.5, green, bold=True)
        add_equation(
            slide,
            formulas["objective"],
            7.1,
            1.74,
            3.4,
            0.4,
            size=17,
            color=ink,
        )

        constraint_rows = [
            (
                2.52,
                "1   Every flow runs exactly once",
                "assign",
                blue,
                14,
            ),
            (
                3.32,
                "2   At most one peer per rank and batch",
                "endpoint",
                purple,
                13,
            ),
            (
                4.12,
                "3   Activate used batches; no gaps",
                "activation",
                green,
                14,
            ),
            (
                4.92,
                "4   Migration load on each link",
                "load",
                orange,
                14,
            ),
            (
                5.72,
                "5   Inference-time lower bound",
                "slowdown",
                green,
                12,
            ),
        ]
        for y, label, key, color, size in constraint_rows:
            add_panel(slide, 4.08, y, 8.5, 0.7, color)
            label_box = slide.shapes.add_textbox(
                Inches(4.28), Inches(y + 0.08), Inches(2.9), Inches(0.5)
            )
            set_text(label_box, label, 10.7, color, bold=True)
            add_equation(
                slide,
                formulas[key],
                6.95,
                y + 0.08,
                5.25,
                0.5,
                size=size,
                color=ink,
            )


def remove_legacy_ilp_math(path: Path) -> None:
    """Remove superseded compatibility equations from the titled ILP slide."""
    with tempfile.TemporaryDirectory() as temp_dir:
        output = Path(temp_dir) / path.name
        with (
            zipfile.ZipFile(path, "r") as source,
            zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
        ):
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename.startswith(
                    "ppt/slides/slide"
                ) and item.filename.endswith(".xml"):
                    root = etree.fromstring(data)
                    namespaces = {
                        "a": ("http://schemas.openxmlformats.org/drawingml/2006/main"),
                        "m": (
                            "http://schemas.openxmlformats.org/officeDocument/2006/math"
                        ),
                        "mc": (
                            "http://schemas.openxmlformats.org/markup-"
                            "compatibility/2006"
                        ),
                    }
                    titles = root.xpath(".//a:t/text()", namespaces=namespaces)
                    if "ILP Formulation" not in titles:
                        target.writestr(item, data)
                        continue
                    for alternate in root.xpath(
                        ".//mc:AlternateContent[.//m:oMath]",
                        namespaces=namespaces,
                    ):
                        alternate.getparent().remove(alternate)
                    substitutions = {
                        "\u200b": "",
                        "\u2001": " ",
                        "\u2004": " ",
                        "\u2005": " ",
                        "\u2009": " ",
                    }
                    for text_node in root.xpath(".//m:t", namespaces=namespaces):
                        if text_node.text is None:
                            continue
                        for old, new in substitutions.items():
                            text_node.text = text_node.text.replace(old, new)
                    data = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                target.writestr(item, data)
        path.write_bytes(output.read_bytes())


def add_card(slide, x: float, title: str, body: str, accent: RGBColor) -> None:
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(2.75),
        Inches(3.45),
        Inches(2.45),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = accent
    card.line.width = Pt(1.6)

    title_box = slide.shapes.add_textbox(
        Inches(x + 0.22), Inches(2.95), Inches(3.0), Inches(0.42)
    )
    set_text(title_box, title, 15, accent, bold=True)

    body_box = slide.shapes.add_textbox(
        Inches(x + 0.22), Inches(3.42), Inches(3.0), Inches(1.48)
    )
    set_text(body_box, body, 13.5, RGBColor(31, 48, 70))


def rebuild_profiling_method(prs: Presentation) -> None:
    """Explain how scheduler cost is compared with clean serving benefit."""
    for slide in prs.slides:
        if not any(
            shape.has_text_frame and shape.text.strip() == "Profiling Method"
            for shape in slide.shapes
        ):
            continue
        for shape in reversed(list(slide.shapes)[4:]):
            remove_shape(shape)

        navy = RGBColor(20, 52, 91)
        blue = RGBColor(47, 111, 174)
        purple = RGBColor(125, 72, 161)
        orange = RGBColor(230, 126, 34)
        green = RGBColor(19, 143, 95)

        question = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.7),
            Inches(1.68),
            Inches(12.0),
            Inches(0.72),
        )
        question.fill.solid()
        question.fill.fore_color.rgb = RGBColor(232, 241, 250)
        question.line.color.rgb = blue
        question.line.width = Pt(1.4)
        set_text(
            question,
            "PURPOSE   Is the added scheduler cost small compared with "
            "the serving time saved?",
            17,
            navy,
            bold=True,
            align=PP_ALIGN.CENTER,
        )

        add_card(
            slide,
            0.7,
            "1   PROFILE COST",
            "Add an NVTX range around the cycle-wide greedy scheduler.\n\n"
            "Nsight Systems records every call on all four Ray workers.",
            purple,
        )
        add_card(
            slide,
            4.94,
            "2   AGGREGATE",
            "Each call builds schedules for all 48 MoE layers.\n\n"
            "Sum call time per rank and use the largest rank total.",
            orange,
        )
        add_card(
            slide,
            9.18,
            "3   MEASURE BENEFIT",
            "Run a clean async E2E A/B with identical settings and no profiler.\n\n"
            "Saved time = batching OFF − batching ON.",
            green,
        )

        for x in (4.32, 8.56):
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.CHEVRON,
                Inches(x),
                Inches(3.62),
                Inches(0.42),
                Inches(0.55),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = blue
            arrow.line.fill.background()

        range_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.95),
            Inches(5.6),
            Inches(11.5),
            Inches(0.76),
        )
        range_box.fill.solid()
        range_box.fill.fore_color.rgb = RGBColor(238, 247, 242)
        range_box.line.color.rgb = green
        range_box.line.width = Pt(1.3)
        set_text(
            range_box,
            'SEARCH IN TRACE   "eplb: schedule migration batches"   '
            "→   compare scheduler cost / serving time saved",
            15,
            navy,
            bold=True,
            align=PP_ALIGN.CENTER,
        )


def replace_ilp_link_image(prs: Presentation) -> None:
    """Refresh the left topology figure while preserving its column height."""
    for slide in prs.slides:
        if not any(
            shape.has_text_frame and shape.text.strip() == "ILP Hotspot Example"
            for shape in slide.shapes
        ):
            continue
        pictures = [
            shape
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        if not pictures:
            return
        target = min(pictures, key=lambda shape: shape.left)
        left, top, width, height = (
            target.left,
            target.top,
            target.width,
            target.height,
        )
        remove_shape(target)
        with Image.open(ILP_LINK_IMAGE) as image:
            new_width = int(height * image.width / image.height)
        new_left = left + (width - new_width) // 2
        slide.shapes.add_picture(
            str(ILP_LINK_IMAGE), new_left, top, width=new_width, height=height
        )
        return


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color) -> None:
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    connector.name = "GreedyGraphArrow"
    connector.line.color.rgb = color
    connector.line.width = Pt(2.2)
    line = connector.line._get_or_add_ln()
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle")
    tail.set("w", "sm")
    tail.set("len", "sm")
    line.append(tail)


def add_click_animations(slide, shape_ids: list[int]) -> None:
    """Add sequential on-click Appear effects for the supplied shapes."""
    old_timing = slide._element.find(qn("p:timing"))
    if old_timing is not None:
        slide._element.remove(old_timing)

    timing = OxmlElement("p:timing")
    timing_list = OxmlElement("p:tnLst")
    root_parallel = OxmlElement("p:par")
    root_time = OxmlElement("p:cTn")
    root_time.set("id", "1")
    root_time.set("dur", "indefinite")
    root_time.set("restart", "never")
    root_time.set("nodeType", "tmRoot")
    root_children = OxmlElement("p:childTnLst")
    sequence = OxmlElement("p:seq")
    sequence.set("concurrent", "1")
    sequence.set("nextAc", "seek")
    main_time = OxmlElement("p:cTn")
    main_time.set("id", "2")
    main_time.set("dur", "indefinite")
    main_time.set("nodeType", "mainSeq")
    main_children = OxmlElement("p:childTnLst")

    next_id = 3
    for shape_id in shape_ids:
        click_parallel = OxmlElement("p:par")
        click_time = OxmlElement("p:cTn")
        click_time.set("id", str(next_id))
        click_time.set("fill", "hold")
        next_id += 1
        click_conditions = OxmlElement("p:stCondLst")
        click_condition = OxmlElement("p:cond")
        click_condition.set("delay", "indefinite")
        click_conditions.append(click_condition)
        click_time.append(click_conditions)

        click_children = OxmlElement("p:childTnLst")
        wrapper_parallel = OxmlElement("p:par")
        wrapper_time = OxmlElement("p:cTn")
        wrapper_time.set("id", str(next_id))
        wrapper_time.set("fill", "hold")
        next_id += 1
        wrapper_conditions = OxmlElement("p:stCondLst")
        wrapper_condition = OxmlElement("p:cond")
        wrapper_condition.set("delay", "0")
        wrapper_conditions.append(wrapper_condition)
        wrapper_time.append(wrapper_conditions)

        wrapper_children = OxmlElement("p:childTnLst")
        effect_parallel = OxmlElement("p:par")
        effect_time = OxmlElement("p:cTn")
        effect_time.set("id", str(next_id))
        effect_time.set("presetID", "1")
        effect_time.set("presetClass", "entr")
        effect_time.set("presetSubtype", "0")
        effect_time.set("fill", "hold")
        effect_time.set("nodeType", "clickEffect")
        next_id += 1
        effect_conditions = OxmlElement("p:stCondLst")
        effect_condition = OxmlElement("p:cond")
        effect_condition.set("delay", "0")
        effect_conditions.append(effect_condition)
        effect_time.append(effect_conditions)

        effect_children = OxmlElement("p:childTnLst")
        visibility = OxmlElement("p:set")
        behavior = OxmlElement("p:cBhvr")
        behavior_time = OxmlElement("p:cTn")
        behavior_time.set("id", str(next_id))
        behavior_time.set("dur", "1")
        behavior_time.set("fill", "hold")
        next_id += 1
        behavior_conditions = OxmlElement("p:stCondLst")
        behavior_condition = OxmlElement("p:cond")
        behavior_condition.set("delay", "0")
        behavior_conditions.append(behavior_condition)
        behavior_time.append(behavior_conditions)
        behavior.append(behavior_time)
        target = OxmlElement("p:tgtEl")
        shape_target = OxmlElement("p:spTgt")
        shape_target.set("spid", str(shape_id))
        target.append(shape_target)
        behavior.append(target)
        attributes = OxmlElement("p:attrNameLst")
        attribute = OxmlElement("p:attrName")
        attribute.text = "style.visibility"
        attributes.append(attribute)
        behavior.append(attributes)
        visibility.append(behavior)
        destination = OxmlElement("p:to")
        visible = OxmlElement("p:strVal")
        visible.set("val", "visible")
        destination.append(visible)
        visibility.append(destination)
        effect_children.append(visibility)
        effect_time.append(effect_children)
        effect_parallel.append(effect_time)
        wrapper_children.append(effect_parallel)
        wrapper_time.append(wrapper_children)
        wrapper_parallel.append(wrapper_time)
        click_children.append(wrapper_parallel)
        click_time.append(click_children)
        click_parallel.append(click_time)
        main_children.append(click_parallel)

    main_time.append(main_children)
    sequence.append(main_time)
    for event_name, list_name in (
        ("onPrev", "p:prevCondLst"),
        ("onNext", "p:nextCondLst"),
    ):
        condition_list = OxmlElement(list_name)
        condition = OxmlElement("p:cond")
        condition.set("evt", event_name)
        condition.set("delay", "0")
        target = OxmlElement("p:tgtEl")
        target.append(OxmlElement("p:sldTgt"))
        condition.append(target)
        condition_list.append(condition)
        sequence.append(condition_list)
    root_children.append(sequence)
    root_time.append(root_children)
    root_parallel.append(root_time)
    timing_list.append(root_parallel)
    timing.append(timing_list)
    slide._element.append(timing)


def rebuild_greedy_algorithm(prs: Presentation) -> None:
    """Add a directed transfer graph and reveal decision rows on clicks."""
    target_slide = next(
        (
            slide
            for slide in prs.slides
            if any(
                shape.has_text_frame
                and shape.text.strip() == "Greedy Algorithm"
                and shape.top < Inches(1.2)
                for shape in slide.shapes
            )
        ),
        None,
    )
    if target_slide is None:
        return

    # Undo artifacts produced before the ILP slide changed the page order.
    for slide in prs.slides:
        if slide.part.partname == target_slide.part.partname:
            continue
        generated_groups = [
            shape
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP
            and shape.name.startswith("DecisionRow")
        ]
        generated_graph = [
            shape for shape in slide.shapes if shape.name.startswith("GreedyGraph")
        ]
        if not generated_groups and not generated_graph:
            continue
        for group in reversed(generated_groups):
            ungroup_shape(group)
        for shape in reversed(generated_graph):
            remove_shape(shape)
        timing = slide._element.find(qn("p:timing"))
        if timing is not None:
            slide._element.remove(timing)

    for slide in (target_slide,):
        if not any(
            shape.has_text_frame
            and shape.text.strip() == "Greedy Algorithm"
            and shape.top < Inches(1.2)
            for shape in slide.shapes
        ):
            continue

        for shape in reversed(list(slide.shapes)):
            if shape.name.startswith("GreedyGraph"):
                remove_shape(shape)
                continue
            if (
                shape.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER
                and shape.left >= Inches(7.0)
                and shape.top >= Inches(4.75)
            ):
                remove_shape(shape)

        decision_groups = [
            shape
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP
            and shape.name.startswith("DecisionRow")
        ]
        if len(decision_groups) != 4:
            decision_groups = []
            for row_index, row_top in enumerate((2.36, 2.97, 3.58, 4.19)):
                row_shapes = [
                    shape
                    for shape in list(slide.shapes)
                    if shape.left >= Inches(5.1)
                    and Inches(row_top - 0.04) <= shape.top < Inches(row_top + 0.57)
                ]
                group = slide.shapes.add_group_shape(row_shapes)
                group.name = f"DecisionRow{row_index}"
                decision_groups.append(group)

        navy = RGBColor(20, 52, 91)
        blue = RGBColor(47, 111, 174)
        light_blue = RGBColor(232, 241, 250)

        graph_title = slide.shapes.add_textbox(
            Inches(5.22), Inches(4.84), Inches(2.6), Inches(0.3)
        )
        graph_title.name = "GreedyGraphTitle"
        set_text(graph_title, "DIRECTED EXPERT TRANSFERS", 11.5, blue, bold=True)

        nodes = {
            0: (7.05, 5.13),
            1: (10.0, 5.13),
            2: (10.0, 5.98),
            3: (7.05, 5.98),
        }
        for rank, (x, y) in nodes.items():
            node = slide.shapes.add_shape(
                MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.5), Inches(0.5)
            )
            node.name = f"GreedyGraphNode{rank}"
            node.fill.solid()
            node.fill.fore_color.rgb = light_blue
            node.line.color.rgb = blue
            node.line.width = Pt(1.5)
            set_text(node, str(rank), 13, navy, bold=True, align=PP_ALIGN.CENTER)

        add_arrow(slide, 7.55, 5.38, 10.0, 5.38, blue)
        add_arrow(slide, 10.25, 5.63, 10.25, 5.98, blue)
        add_arrow(slide, 10.0, 6.23, 7.55, 6.23, blue)
        add_arrow(slide, 7.3, 5.63, 7.3, 5.98, blue)

        labels = [
            ("F₀  {E₁,E₇}", 8.2, 5.08, 1.3),
            ("F₂  {E₄}", 10.38, 5.62, 1.05),
            ("F₁  {E₃}", 8.32, 6.28, 1.05),
            ("F₃  {E₉}", 6.05, 5.62, 1.05),
        ]
        for index, (text, x, y, width) in enumerate(labels):
            label = slide.shapes.add_textbox(
                Inches(x), Inches(y), Inches(width), Inches(0.25)
            )
            label.name = f"GreedyGraphLabel{index}"
            set_text(label, text, 10.5, navy, bold=True, align=PP_ALIGN.CENTER)

        decision_groups.sort(key=lambda shape: shape.top)
        add_click_animations(slide, [shape.shape_id for shape in decision_groups])
        return


def enlarge_small_text(prs: Presentation) -> None:
    """Raise small labels while respecting compact tables and diagrams."""
    for slide_number, slide in enumerate(prs.slides, start=1):
        is_ilp_formulation = any(
            shape.has_text_frame and shape.text.strip() == "ILP Formulation"
            for shape in slide.shapes
        )
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text.strip():
                continue

            if is_ilp_formulation:
                continue

            y = shape.top / Inches(1)
            explicit_sizes = [
                run.font.size.pt
                for paragraph in shape.text_frame.paragraphs
                for run in paragraph.runs
                if run.font.size is not None
            ]

            # Small section labels above the page title.
            if y < 0.35:
                if explicit_sizes and min(explicit_sizes) < 10:
                    set_shape_font_size(shape, 10)
                continue

            # The compact navigation strip on the setup slide is decorative.
            if slide_number == 10 and y > 6.3:
                if explicit_sizes and min(explicit_sizes) < 9:
                    set_shape_font_size(shape, 9)
                continue

            minimum = {
                1: 11,
                3: 11.5,
                5: 11.5,
                6: 12,
                8: 12,
                9: 11.5,
                10: 11.5,
                11: 11.5,
                12: 11,
                13: 11,
                14: 11,
                15: 11,
                16: 12,
            }.get(slide_number, 11)

            if explicit_sizes and min(explicit_sizes) < minimum:
                set_shape_font_size(shape, minimum)


def main() -> None:
    prs = Presentation(DECK)
    removed = remove_page_subtitles(prs)
    rebuild_ilp_formulation(prs)
    rebuild_profiling_method(prs)
    rebuild_greedy_algorithm(prs)
    replace_ilp_link_image(prs)
    enlarge_small_text(prs)
    prs.save(DECK)
    remove_legacy_ilp_math(DECK)
    print(f"Removed {removed} page subtitles and enlarged small text.")


if __name__ == "__main__":
    main()
