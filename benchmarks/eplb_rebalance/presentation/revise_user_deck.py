from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
DECK = HERE / "eplb_migration_batching_pr52641.pptx"
CODE_SCREENSHOT = HERE / "github_first_fit_code.png"
PR_URL = "https://github.com/vllm-project/vllm/pull/52641"

BG = RGBColor(247, 249, 252)
PANEL = RGBColor(255, 255, 255)
PANEL_2 = RGBColor(237, 243, 250)
INK = RGBColor(22, 34, 54)
MUTED = RGBColor(89, 108, 133)
CYAN = RGBColor(13, 138, 166)
BLUE = RGBColor(51, 102, 214)
GREEN = RGBColor(19, 155, 103)
MAGENTA = RGBColor(162, 75, 196)
ORANGE = RGBColor(215, 122, 11)
GRID = RGBColor(216, 226, 239)
WHITE = RGBColor(255, 255, 255)


prs = Presentation(DECK)
blank = prs.slide_layouts[6]


def slide_text(slide) -> str:
    return "\n".join(
        shape.text.strip()
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.text.strip()
    )


def delete_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def add_bg(slide) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()


def add_box(slide, x, y, w, h, *, fill=PANEL, line=GRID, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        kind, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=20,
    color=INK,
    bold=False,
    font="Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    url=None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.04)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if url:
        run.hyperlink.address = url
    return box


def add_title(slide, section, title, subtitle=None):
    add_text(slide, section.upper(), 0.55, 0.23, 2.8, 0.28,
             size=10, color=CYAN, bold=True)
    add_text(slide, title, 0.55, 0.56, 12.0, 0.53,
             size=27, color=INK, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.08, 12.0, 0.35,
                 size=12, color=MUTED)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(1.43),
        Inches(12.17), Inches(0.012)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = GRID
    line.line.fill.background()


def add_image_fit(slide, path, x, y, w, h, *, crop=False):
    with Image.open(path) as image:
        iw, ih = image.size
    image_ratio = iw / ih
    box_ratio = w / h
    if crop:
        picture = slide.shapes.add_picture(
            str(path), Inches(x), Inches(y), Inches(w), Inches(h)
        )
        if image_ratio > box_ratio:
            shown = box_ratio / image_ratio
            picture.crop_left = picture.crop_right = (1 - shown) / 2
        else:
            shown = image_ratio / box_ratio
            picture.crop_top = picture.crop_bottom = (1 - shown) / 2
        return picture
    if image_ratio > box_ratio:
        fitted_w = w
        fitted_h = w / image_ratio
        x += 0
        y += (h - fitted_h) / 2
    else:
        fitted_h = h
        fitted_w = h * image_ratio
        x += (w - fitted_w) / 2
    return slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), Inches(fitted_w), Inches(fitted_h)
    )


def move_slide(slide, index):
    slide_ids = prs.slides._sldIdLst
    slide_id = slide_ids[prs.slides.index(slide)]
    slide_ids.remove(slide_id)
    slide_ids.insert(index, slide_id)


def make_flow(slide, text, x, y, color):
    add_box(slide, x, y, 2.22, 0.42, fill=PANEL_2, line=color)
    add_text(slide, text, x + 0.07, y + 0.05, 2.08, 0.29,
             size=11, color=INK, bold=True, align=PP_ALIGN.CENTER)


for slide in prs.slides:
    for shape in list(slide.shapes):
        text = shape.text.strip() if hasattr(shape, "text") else ""
        is_source = text.lower().startswith("source:") or text.lower().startswith(
            "sources:"
        )
        is_page_number = (
            shape.top >= Inches(6.9) and re.fullmatch(r"\d{2}", text) is not None
        )
        if is_source or is_page_number:
            delete_shape(shape)


if any("Greedy First-Fit Scheduling" in slide_text(s) for s in prs.slides):
    raise RuntimeError("The requested revision has already been applied.")


hardware_index = next(
    i for i, slide in enumerate(prs.slides) if "Hardware" in slide_text(slide)
)
metric_slide = next(
    slide for slide in prs.slides if "Metric Results" in slide_text(slide)
)


# New slide: concrete greedy example after Hardware.
algorithm_slide = prs.slides.add_slide(blank)
add_bg(algorithm_slide)
add_title(
    algorithm_slide,
    "Part II · E2E",
    "Greedy First-Fit Scheduling",
    "Coalesce equal rank pairs, then place each flow in the first compatible batch.",
)
add_box(algorithm_slide, 0.62, 1.72, 5.55, 4.98, fill=PANEL, line=CYAN)
add_text(algorithm_slide, "Example input flows", 0.9, 1.94, 2.1, 0.3,
         size=13, color=CYAN, bold=True)
make_flow(algorithm_slide, "F₀: 0 → 1  [E1, E7]", 0.9, 2.37, BLUE)
make_flow(algorithm_slide, "F₁: 2 → 3  [E3]", 3.48, 2.37, GREEN)
make_flow(algorithm_slide, "F₂: 1 → 2  [E4]", 0.9, 2.91, MAGENTA)
make_flow(algorithm_slide, "F₃: 0 → 3  [E9]", 3.48, 2.91, ORANGE)
add_text(algorithm_slide, "FIRST FIT", 2.47, 3.52, 1.85, 0.28,
         size=11, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
add_text(algorithm_slide, "↓", 3.15, 3.73, 0.45, 0.35,
         size=24, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
add_box(algorithm_slide, 0.9, 4.18, 4.98, 0.83, fill=PANEL_2, line=GREEN)
add_text(algorithm_slide, "Batch 0", 1.08, 4.28, 0.9, 0.26,
         size=12, color=GREEN, bold=True)
add_text(algorithm_slide, "F₀: 0→1     F₁: 2→3", 2.05, 4.27, 3.45, 0.3,
         size=14, color=INK, bold=True, align=PP_ALIGN.CENTER)
add_text(algorithm_slide, "All four endpoints are distinct", 1.1, 4.62, 4.55, 0.22,
         size=9, color=MUTED, align=PP_ALIGN.CENTER)
add_box(algorithm_slide, 0.9, 5.18, 4.98, 0.83, fill=PANEL_2, line=MAGENTA)
add_text(algorithm_slide, "Batch 1", 1.08, 5.28, 0.9, 0.26,
         size=12, color=MAGENTA, bold=True)
add_text(algorithm_slide, "F₂: 1→2     F₃: 0→3", 2.05, 5.27, 3.45, 0.3,
         size=14, color=INK, bold=True, align=PP_ALIGN.CENTER)
add_text(algorithm_slide, "Conflicting flows move to the next batch", 1.1, 5.62, 4.55, 0.22,
         size=9, color=MUTED, align=PP_ALIGN.CENTER)
add_box(algorithm_slide, 6.45, 1.72, 6.25, 4.98, fill=PANEL, line=GRID)
add_image_fit(
    algorithm_slide, CODE_SCREENSHOT, 6.57, 1.84, 6.01, 4.38, crop=True
)
add_text(
    algorithm_slide,
    "Deterministic first-fit edge coloring: each rank has at most one peer per batch.",
    6.76,
    6.28,
    5.62,
    0.28,
    size=10,
    color=CYAN,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# New slide: exact benchmark command before Metric Results.
command_slide = prs.slides.add_slide(blank)
add_bg(command_slide)
add_title(
    command_slide,
    "Part II · E2E",
    "Benchmark Command",
    "The off/on runs use the same request stream and a fresh server.",
)
add_box(command_slide, 0.68, 1.72, 11.98, 4.93, fill=PANEL, line=BLUE)
command = '''vllm bench serve --backend vllm --model "$MODEL" --port 8000 \\
  --dataset-name random --random-prefix-len 300 \\
  --random-input-len 200 --random-output-len 300 \\
  --num-prompts 500 --max-concurrency 32 --seed 0 --temperature 0 \\
  --percentile-metrics ttft,tpot,e2el --metric-percentiles 50,99 \\
  --ignore-eos --disable-tqdm --save-result \\
  --result-dir "$RESULTS" --result-filename bench.json'''
add_text(
    command_slide,
    command,
    1.0,
    2.02,
    11.35,
    3.55,
    size=16,
    color=INK,
    font="Consolas",
    valign=MSO_ANCHOR.TOP,
)
add_box(command_slide, 1.0, 5.78, 11.35, 0.52, fill=PANEL_2, line=GREEN)
add_text(
    command_slide,
    "500 prompts  ·  concurrency 32  ·  deterministic decoding  ·  TTFT / TPOT / E2EL P50 and P99",
    1.2,
    5.89,
    10.95,
    0.27,
    size=11,
    color=GREEN,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# Extend the existing user-edited Metric Results slide with the PR table.
picture = next(shape for shape in metric_slide.shapes if shape.shape_type == 13)
picture.left = Inches(3.82)
picture.top = Inches(1.57)
picture.width = Inches(5.70)
picture.height = Inches(2.38)
for shape in metric_slide.shapes:
    if shape.shape_type == 1 and shape.top > Inches(1.5):
        shape.left = Inches(0.56)
        shape.top = Inches(1.53)
        shape.width = Inches(12.2)
        shape.height = Inches(2.48)
        break

rows = [
    [
        "Batching", "Duration\n(s)", "Request/s", "Output\ntok/s",
        "TTFT P50 / P99 (ms)", "TPOT P50 / P99 (ms)",
        "E2EL P50 / P99 (ms)",
    ],
    [
        "off", "4,756.10", "0.1051", "31.54",
        "49,180.71 / 75,740.25", "828.53 / 1,007.82",
        "303,462.69 / 322,184.88",
    ],
    [
        "on", "4,650.17", "0.1075", "32.26",
        "48,599.04 / 74,756.17", "812.44 / 999.01",
        "293,821.76 / 317,351.61",
    ],
]
table_shape = metric_slide.shapes.add_table(
    3, 7, Inches(0.56), Inches(4.22), Inches(12.2), Inches(2.05)
)
table = table_shape.table
widths = [0.78, 1.1, 1.0, 1.05, 2.38, 2.32, 2.84]
for col, width in zip(table.columns, widths):
    col.width = Inches(width)
table.rows[0].height = Inches(0.63)
table.rows[1].height = Inches(0.71)
table.rows[2].height = Inches(0.71)
for row_idx, values in enumerate(rows):
    for col_idx, value in enumerate(values):
        cell = table.cell(row_idx, col_idx)
        cell.text = value
        cell.margin_left = cell.margin_right = Inches(0.04)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if row_idx == 0:
            cell.fill.fore_color.rgb = PANEL_2
        elif row_idx == 2:
            cell.fill.fore_color.rgb = RGBColor(236, 248, 242)
        else:
            cell.fill.fore_color.rgb = WHITE
        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8 if row_idx == 0 else 9)
                run.font.bold = row_idx == 0 or col_idx == 0
                run.font.color.rgb = GREEN if row_idx == 2 else INK


# New final slide.
thanks_slide = prs.slides.add_slide(blank)
add_bg(thanks_slide)
add_text(thanks_slide, "vLLM · EPLB", 0.72, 0.72, 3.2, 0.35,
         size=13, color=CYAN, bold=True)
add_text(thanks_slide, "Thank you", 0.72, 1.65, 11.9, 1.05,
         size=48, color=INK, bold=True, align=PP_ALIGN.CENTER)
add_text(thanks_slide, "Questions & Discussion", 0.72, 2.82, 11.9, 0.58,
         size=22, color=MUTED, align=PP_ALIGN.CENTER)
add_box(thanks_slide, 3.18, 4.05, 6.98, 0.72, fill=PANEL_2, line=CYAN)
add_text(
    thanks_slide,
    "github.com/vllm-project/vllm/pull/52641",
    3.38,
    4.2,
    6.58,
    0.34,
    size=14,
    color=CYAN,
    bold=True,
    align=PP_ALIGN.CENTER,
    url=PR_URL,
)
for i, color in enumerate((BLUE, CYAN, GREEN, MAGENTA)):
    add_box(thanks_slide, 4.18 + i * 1.35, 5.62, 0.9, 0.12,
            fill=color, line=color, radius=False)


move_slide(algorithm_slide, hardware_index + 1)
move_slide(command_slide, hardware_index + 2)

tmp = DECK.with_suffix(".tmp.pptx")
prs.save(tmp)
tmp.replace(DECK)
print(DECK)
